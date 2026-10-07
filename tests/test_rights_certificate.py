import sys
import tempfile
import unittest
import hashlib
from pathlib import Path
from unittest.mock import MagicMock, patch

from bs4 import BeautifulSoup
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches

BACKEND_PATH = Path(__file__).resolve().parents[1] / "automation-service" / "backend"
if str(BACKEND_PATH) not in sys.path:
    sys.path.insert(0, str(BACKEND_PATH))

from app.services.rights_certificate import (  # noqa: E402
    _extract_management_fee,
    _extract_management_fee_amount,
    _extract_dividend_requests,
    _extract_myungseung_rights_analysis,
    _extract_rights,
    _clean_court_label,
    _paginate_narrative_sections,
    _right_dividend_request_status,
    _tenant_ocr_text_indicates_no_tenants,
    _is_no_tenant_record,
    _is_name_like_token,
    _looks_like_person_name,
    _guess_sale_spec_tenant_name,
    _sale_spec_name_before_source,
    _parse_sale_spec_occupancy_blocks,
    NO_TENANTS_TEXT,
    build_tenant_analysis_text,
    analyze_surplus,
    analyze_registered_takeover_rights,
    analyze_tenants,
    build_template_data,
    calculate_surplus_basis,
    export_pptx_to_pdf,
    extract_rights_context,
    extract_rights_context_by_ocr,
    extract_sale_spec_tenant_context_by_ocr,
    render_certificate_template,
    render_certificate_pptx_template,
    merge_rights,
    parse_rights_from_ocr,
    parse_sale_spec_tenants_from_pdf_text,
    parse_sale_spec_tenants_from_ocr,
    parse_tenants_from_ocr,
    registered_right_amount_total,
    substantive_registered_rights,
)
from app.services.rights_checklist import State, build_checklist_from_pipeline  # noqa: E402
from app.services.briefing_rights import _build_surplus_description as build_briefing_surplus_description  # noqa: E402
from app.services import briefing_cost_images, briefing_opinion, briefing_rights, capturer, crawler, forced_execution_estimator, ppt_builder  # noqa: E402


def _slide_text(slide) -> str:
    values = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            values.append(shape.text_frame.text)
        if getattr(shape, "has_table", False):
            values.extend(cell.text for row in shape.table.rows for cell in row.cells)
    return "\n".join(values)


def _slide_image_hashes(slide) -> set[str]:
    return {
        hashlib.sha256(rel.target_part.blob).hexdigest()
        for rel in slide.part.rels.values()
        if not rel.is_external and rel.reltype.endswith("/image")
    }


class RightsCertificateTests(unittest.TestCase):
    def test_briefing_cost_images_use_bid_three_for_acquisition_tax_and_fixed_loan_priority(self):
        ctx = briefing_cost_images.build_cost_context(
            {
                "market_price": 620_000_000,
                "bid_price_1": 590_000_000,
                "bid_price_2": 580_000_000,
                "bid_price_3": 570_000_000,
                "difference_amount": 5_000_000,
                "property_tax_type": "house",
                "house_count": "1",
                "regulated_area": False,
                "unpaid_management_fee": 300_000,
                "service_fee_basis": "appraised",
                "service_fee_rate": 1,
                "fixed_loan_amount": 400_000_000,
                "loan_base_amount": 700_000_000,
                "ltv_limit": 80,
                "bid_price_loan_limit": 80,
                "loan_room_deduction": 50_000_000,
                "bank_loan_note": "한도 4억원 우선",
            },
            {
                "appraised_price": "668,000,000원",
                "min_price": "467,600,000원",
                "building_area_m2": "40",
                "eviction_cost_values": {"flat_total": 4_064_000},
            },
        )

        self.assertEqual(ctx["tax1"]["price"], 570_000_000)
        self.assertEqual(ctx["scenarios"][0]["loan"], 400_000_000)
        self.assertEqual(ctx["scenarios"][1]["loan"], 400_000_000)
        self.assertEqual(ctx["scenarios"][2]["loan"], 400_000_000)
        self.assertEqual(ctx["scenarios"][0]["consulting_fee"], 6_680_000)
        self.assertEqual(ctx["scenarios"][0]["eviction_cost"], 4_064_000)
        self.assertEqual(ctx["scenarios"][0]["unpaid_management_fee"], 300_000)
        self.assertEqual(ctx["tax1"]["total_tax"], 6_270_000)
        self.assertEqual(
            [(row["label"], row["bid"]) for row in ctx["bid_rows"]],
            [
                ("감정가", 668_000_000),
                ("...", 605_000_000),
                ("...", 600_000_000),
                ("...", 595_000_000),
                ("낙찰 우위입찰가", 590_000_000),
                ("경쟁 균형입찰가", 580_000_000),
                ("안정 투자입찰가", 570_000_000),
                ("...", 565_000_000),
                ("최저 입찰가", 467_600_000),
            ],
        )

    def test_briefing_cost_images_clamp_minimum_price_when_it_exceeds_appraisal(self):
        ctx = briefing_cost_images.build_cost_context(
            {
                "bid_price_1": 1_820_000_000,
                "bid_price_2": 1_810_000_000,
                "bid_price_3": 1_800_000_000,
            },
            {"appraised_price": 1_847_000_000, "min_price": 3_030_000_000},
        )

        self.assertEqual(ctx["appraised"], 1_847_000_000)
        self.assertEqual(ctx["minimum"], 1_847_000_000)

    def test_briefing_cost_images_blank_fixed_loan_uses_lower_of_appraisal_and_bid_ltv(self):
        ctx = briefing_cost_images.build_cost_context(
            {
                "market_price": 700_000_000,
                "bid_price_1": 590_000_000,
                "bid_price_2": 580_000_000,
                "bid_price_3": 570_000_000,
                "fixed_loan_amount": "",
                "loan_base_amount": "",
                "ltv_limit": 70,
                "bid_price_loan_limit": 80,
            },
            {"appraised_price": 668_000_000, "min_price": 467_600_000},
        )

        self.assertEqual(ctx["scenarios"][0]["loan"], 467_600_000)
        self.assertEqual(ctx["scenarios"][1]["loan"], 464_000_000)
        self.assertEqual(ctx["scenarios"][2]["loan"], 456_000_000)
        self.assertIn("감정가 70", ctx["scenarios"][0]["loan_note"])
        self.assertNotIn("고정값", ctx["scenarios"][0]["loan_note"])

    def test_briefing_cost_images_default_bank_loan_note_matches_planner_policy(self):
        ctx = briefing_cost_images.build_cost_context(
            {
                "bid_price_1": 590_000_000,
                "bid_price_2": 580_000_000,
                "bid_price_3": 570_000_000,
                "fixed_loan_amount": "",
                "loan_base_amount": "",
                "ltv_limit": 40,
                "bid_price_loan_limit": 80,
                "bank_loan_note": "",
            },
            {"appraised_price": 668_000_000, "min_price": 467_600_000},
        )

        self.assertEqual(ctx["scenarios"][0]["loan"], 267_200_000)
        self.assertEqual(ctx["scenarios"][0]["loan_note"], "감정가 40%,낙찰가80% 중 낮은금액으로 대출이 가능합니다.")

    def test_briefing_cost_images_custom_bank_loan_note_is_preserved(self):
        ctx = briefing_cost_images.build_cost_context(
            {
                "bid_price_1": 590_000_000,
                "bid_price_2": 580_000_000,
                "bid_price_3": 570_000_000,
                "ltv_limit": 40,
                "bid_price_loan_limit": 80,
                "bank_loan_note": "은행 확인 후 별도 적용",
            },
            {"appraised_price": 668_000_000, "min_price": 467_600_000},
        )

        self.assertEqual(ctx["scenarios"][0]["loan_note"], "은행 확인 후 별도 적용")

    def test_briefing_cost_images_bid_and_cost_tables_use_bold_text(self):
        ctx = briefing_cost_images.build_cost_context(
            {
                "market_price": 620_000_000,
                "bid_price_1": 590_000_000,
                "bid_price_2": 580_000_000,
                "bid_price_3": 570_000_000,
                "service_fee_rate": 1,
                "ltv_limit": 80,
                "bid_price_loan_limit": 80,
            },
            {
                "appraised_price": 668_000_000,
                "min_price": 467_600_000,
                "building_area_m2": 40,
                "eviction_cost_values": {"flat_total": 4_064_000},
            },
        )
        bold_fonts = (
            briefing_cost_images.FONT_18_B,
            briefing_cost_images.FONT_20_B,
            briefing_cost_images.FONT_22_B,
            briefing_cost_images.FONT_24_B,
            briefing_cost_images.FONT_26_B,
            briefing_cost_images.FONT_34_B,
        )
        used_fonts: list[object] = []

        def capture_text(_draw, _xy, _text, font, *args, **kwargs):
            used_fonts.append(font)

        with tempfile.TemporaryDirectory() as tmp, patch.object(briefing_cost_images, "_draw_text", side_effect=capture_text):
            briefing_cost_images.render_bid_price_image(ctx, Path(tmp) / "bid.png")
            briefing_cost_images.render_acquisition_cost_sheet_image(ctx, Path(tmp) / "cost.png")

        self.assertTrue(used_fonts)
        self.assertTrue(all(any(font is bold for bold in bold_fonts) for font in used_fonts))

    def test_briefing_cost_images_render_three_pngs(self):
        with tempfile.TemporaryDirectory() as tmp:
            rendered = briefing_cost_images.render_briefing_cost_images(
                {
                    "market_price": 620_000_000,
                    "bid_price_1": 590_000_000,
                    "bid_price_2": 580_000_000,
                    "bid_price_3": 570_000_000,
                    "property_tax_type": "house",
                    "house_count": "1",
                    "regulated_area": False,
                    "unpaid_management_fee": 300_000,
                    "service_fee_rate": 1,
                    "fixed_loan_amount": 400_000_000,
                },
                {
                    "appraised_price": 668_000_000,
                    "min_price": 467_600_000,
                    "building_area_m2": 40,
                    "eviction_cost_values": {"flat_total": 4_064_000},
                },
                tmp,
            )

            self.assertEqual(set(rendered.keys()), {"acquisition-tax", "loan-bid-estimator", "acquisition-cost-sheet"})
            for path in rendered.values():
                self.assertTrue(Path(path).exists())
                with Image.open(path) as img:
                    self.assertGreaterEqual(img.width, 1400)
                    self.assertGreaterEqual(img.height, 800)

    def test_eviction_fixed_total_uses_myauction_detail_main_price(self):
        class FakeDriver:
            def execute_script(self, _script):
                return "총 407만원"

        flat_price = capturer._extract_eviction_flat_rate_main_price(FakeDriver())
        self.assertEqual(flat_price, 4_070_000)

        values = forced_execution_estimator.build_eviction_cost_values({
            "item_type": "아파트",
            "myauction_eviction_costs": {
                "filingFee": 150_000,
                "transportStorage": 2_200_000,
                "laborTotal": 2_080_000,
                "laborWorkers": 16,
                "locksmith": 200_000,
                "ladderTruck": 350_000,
                "witness": 100_000,
                "grandTotal": 5_080_000,
                "flatRateMainPrice": flat_price,
            },
        })
        self.assertEqual(values["flat_total"], 4_070_000)
        self.assertEqual(values["명도_정액제총액"], 4_070_000)
        self.assertEqual(values["명도_총명도비용"], 5_080_000)
        self.assertEqual(values["normal_execution_cost"], 5_080_000)
        self.assertEqual(values["eviction_flat_total_source"], "myauction_detail")
        self.assertNotEqual(values["flat_total"], round((values["attorney_fee"] + 5_080_000) * 0.8))

    def test_eviction_cost_basis_capture_replaces_left_estimate_table(self):
        prs = Presentation(str(BACKEND_PATH / "templates" / "sample2_configured.pptx"))
        slide = ppt_builder.find_slide_by_note_key(prs, "FORCED_EXECUTION_ESTIMATE_BOX")
        self.assertIsNotNone(slide)
        before_pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]

        with tempfile.TemporaryDirectory() as tmp:
            image_path = Path(tmp) / "eviction-cost-basis.png"
            Image.new("RGB", (727, 884), "white").save(image_path)
            self.assertTrue(ppt_builder.insert_eviction_cost_basis_image(prs, str(image_path)))

        after_pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(after_pictures), len(before_pictures) + 1)
        inserted = after_pictures[-1]
        self.assertEqual(inserted.left, 506628)
        self.assertEqual(inserted.top, 1606378)
        self.assertEqual(inserted.width, 4534930)
        self.assertEqual(inserted.height, 5276335)

    def test_eviction_cost_basis_uses_user_yellow_box_size_when_present(self):
        prs = Presentation(str(BACKEND_PATH / "templates" / "sample2_configured.pptx"))
        slide = ppt_builder.find_slide_by_note_key(prs, "FORCED_EXECUTION_ESTIMATE_BOX")
        self.assertIsNotNone(slide)
        marker = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.73),
            Inches(1.42),
            Inches(4.82),
            Inches(5.66),
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor(255, 255, 0)

        with tempfile.TemporaryDirectory() as tmp:
            image_path = Path(tmp) / "eviction-cost-basis.png"
            Image.new("RGB", (727, 884), "white").save(image_path)
            self.assertTrue(ppt_builder.insert_eviction_cost_basis_image(prs, str(image_path)))

        pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        inserted = pictures[-1]
        self.assertEqual(inserted.left, Inches(0.73))
        self.assertEqual(inserted.top, Inches(1.42))
        self.assertEqual(inserted.width, Inches(4.82))
        self.assertEqual(inserted.height, Inches(5.66))
        self.assertIsNone(ppt_builder.find_explicit_yellow_box(slide))

    def test_briefing_images_are_inserted_with_contained_aspect_ratio(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.2), Inches(0.2), Inches(1), Inches(0.3))
        title.text = "테스트이미지"
        marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(4), Inches(4))
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor(255, 192, 0)

        with tempfile.TemporaryDirectory() as tmp:
            image_path = Path(tmp) / "wide.png"
            Image.new("RGB", (1000, 500), "navy").save(image_path)
            ppt_builder.insert_single_image(prs, "테스트이미지", str(image_path))

        pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        self.assertEqual(len(pictures), 1)
        inserted = pictures[0]
        self.assertEqual(inserted.width, Inches(4))
        self.assertEqual(inserted.height, Inches(2))
        self.assertEqual(inserted.left, Inches(1))
        self.assertEqual(inserted.top, Inches(2))

    def test_rights_analysis_opinion_continues_to_next_slide_for_many_tenants(self):
        prs = Presentation(str(BACKEND_PATH / "templates" / "sample2_configured.pptx"))
        before_count = len(prs.slides)
        tenant_lines = [
            f"- 점유자 성명: 임차인{i} / 점유구분: 주거 / 보증금: {i * 10_000_000:,}원 / "
            f"전입일: 2024.0{i}.01 / 확정일: 2024.0{i}.02 / 배당요구일: 2024.0{i}.03."
            for i in range(1, 8)
        ]
        opinion = "\n".join([
            "1) 말소기준 및 등기부상 소멸사항",
            "- 말소기준권리 이후의 권리는 매각으로 말소됩니다.",
            "",
            "2) 임차권리 인수사항",
            *tenant_lines,
            "",
            "3) 경매취하 / 무잉여 가능성",
            "- 현재 확인된 진행상황을 기준으로 별도 취하 접수는 확인되지 않았습니다.",
        ])

        self.assertTrue(ppt_builder.apply_rights_analysis_opinion(prs, opinion))

        all_text = "\n".join(_slide_text(slide) for slide in prs.slides)
        for i in range(1, 8):
            self.assertIn(f"임차인{i}", all_text)
        self.assertGreater(len(prs.slides), before_count)
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "many-tenants.pptx"
            prs.save(output)
            reloaded = Presentation(str(output))
            reloaded_text = "\n".join(_slide_text(slide) for slide in reloaded.slides)
            self.assertIn("임차인7", reloaded_text)

    def test_optional_empty_briefing_pages_are_removed_but_filled_pages_remain(self):
        prs = Presentation(str(BACKEND_PATH / "templates" / "sample2_configured.pptx"))
        with tempfile.TemporaryDirectory() as tmp:
            image_path = Path(tmp) / "filled.png"
            Image.new("RGB", (640, 360), "blue").save(image_path)
            slide = prs.slides[18]  # 실거래가 선택 페이지
            slide.shapes.add_picture(str(image_path), Inches(1), Inches(1), width=Inches(4), height=Inches(2.25))

        removed = ppt_builder.cleanup_optional_empty_slides(prs)
        notes = [
            (slide.notes_slide.notes_text_frame.text or "")
            for slide in prs.slides
        ]
        self.assertGreaterEqual(removed, 8)
        self.assertTrue(any("실거래가" in _slide_text(slide) for slide in prs.slides))
        self.assertFalse(any("네이버 매물 현황 -2" in _slide_text(slide) for slide in prs.slides))
        self.assertFalse(any("SLIDE_KEY=OPINION_RELATED_LAW" in note for note in notes))

    def test_appraisal_sentence_cleanup_removes_entities_and_mechanical_endings(self):
        sentence = briefing_opinion._ensure_sentence(
            '소재 &amp; 아파트로 이용 중이며, &quot;중로1류&quot;와 접하고 붙임 사진과 같음입니다'
        )
        self.assertNotIn("&amp;", sentence)
        self.assertNotIn("&quot;", sentence)
        self.assertNotIn("같음입니다", sentence)
        self.assertIn("및", sentence)
        self.assertIn("같습니다", sentence)

    def test_customer_style_linter_blocks_internal_markers(self):
        with self.assertRaisesRegex(ValueError, "고객 출력 문체 린터 실패"):
            briefing_opinion.assert_customer_facing_texts({
                "special_opinion": "checklistRows 판정식 YAML {내부 확인}\n[ ] 담당자 체크",
            })

        briefing_opinion.assert_customer_facing_texts({
            "special_opinion": "본건은 매각기일 전 현황조사와 등기사항을 재확인하는 조건으로 검토합니다.",
        })

    def test_case_1232734_court_keeps_parent_and_branch(self):
        self.assertEqual(
            _clean_court_label("수원지방법원 안양지원 경매4계 2022 타경 102285"),
            "수원지방법원 안양지원",
        )
        self.assertEqual(
            _clean_court_label("안양지원 경매4계 2022 타경 102285 [아파트]"),
            "수원지방법원 안양지원",
        )
        # 서부지원처럼 둘 이상의 본원에 속하는 짧은 명칭은 추측하지 않는다.
        self.assertEqual(_clean_court_label("서부지원 경매1계"), "서부지원")

        soup = BeautifulSoup(
            """
            <h2>수원지방법원 관련사건</h2>
            <div id="header_detailz"><h2>안양지원 경매4계 <span class="blue">2022 타경 102285</span> [아파트]</h2></div>
            """,
            "html.parser",
        )
        parsed = crawler.parse_myauction_detail(soup, "https://www.my-auction.co.kr/view/1232734")
        self.assertEqual(parsed["case_number"], "2022 타경 102285")
        self.assertEqual(_clean_court_label(parsed["court"]), "수원지방법원 안양지원")

    def test_auction_date_never_uses_dividend_deadline(self):
        soup = BeautifulSoup(
            """
            <div id="dtl_table">
              <table>
                <tr><th>배당요구종기일</th><td>2026.10.12</td></tr>
                <tr><th>청구금액</th><td>82,859,092원</td></tr>
              </table>
            </div>
            <div class="plan_day"><span class="pink">매각기일 2026.11.03 10:30</span></div>
            """,
            "html.parser",
        )
        parsed = crawler.parse_myauction_detail(soup, "https://www.my-auction.co.kr/view/test")
        self.assertEqual(parsed["auction_date"], "2026.11.03 10:30")
        self.assertEqual(parsed["입찰기일"], "2026.11.03 10:30")
        self.assertNotEqual(parsed["auction_date"], "2026.10.12")

        detail_header = BeautifulSoup(
            """
            <div id="dtl_title">
              <h3>사건 정보</h3>
              <div>기타</div>
              <div>
                <ul>
                  <li><span class="sale_txt"><span>2026.09.29 10:30</span></span></li>
                </ul>
              </div>
            </div>
            <p class="plan_day"><span class="pink">매각기일 2026.09.01 10:30</span></p>
            <div id="dtl_table">
              <table><tr><th>배당요구종기일</th><td>2026.08.01</td></tr></table>
            </div>
            """,
            "html.parser",
        )
        parsed_detail_header = crawler.parse_myauction_detail(detail_header, "https://www.my-auction.co.kr/view/test")
        self.assertEqual(parsed_detail_header["auction_date"], "2026.09.29 10:30")
        self.assertNotEqual(parsed_detail_header["auction_date"], "2026.09.01 10:30")

        detail_header_with_minimum = BeautifulSoup(
            """
            <div id="detail_left">
              <table class="tbl_detail">
                <tr><th>경매종류</th><td>부동산임의경매</td><th>감정가</th><td>1,847,000,000원</td></tr>
                <tr><th>최저가</th><td>3,030,000,000원</td><th>입찰보증금</th><td>303,000,000원</td></tr>
              </table>
            </div>
            <div id="dtl_title">
              <h3>사건 정보</h3>
              <div>기타</div>
              <div>
                <ul>
                  <li><span class="sale_txt"><span>2026.09.29 10:30</span></span> 최저가 1,650,000,000원 입찰보증금 165,000,000원 89%</li>
                </ul>
              </div>
            </div>
            """,
            "html.parser",
        )
        parsed_header_with_minimum = crawler.parse_myauction_detail(detail_header_with_minimum, "https://www.my-auction.co.kr/view/test")
        self.assertEqual(parsed_header_with_minimum["auction_date"], "2026.09.29 10:30")
        self.assertEqual(parsed_header_with_minimum["min_price"], "1,650,000,000원")
        self.assertEqual(parsed_header_with_minimum["deposit"], "165,000,000원")
        self.assertEqual(parsed_header_with_minimum["min_rate"], "89%")

        same_row = BeautifulSoup(
            """
            <div id="dtl_table">
              <table>
                <tr>
                  <th>경매종류</th><td>부동산임의경매</td>
                  <th>매각기일</th><td>2026.09.08</td>
                  <th>배당요구종기일</th><td>2026.08.01</td>
                </tr>
              </table>
            </div>
            """,
            "html.parser",
        )
        parsed_same_row = crawler.parse_myauction_detail(same_row, "https://www.my-auction.co.kr/view/test")
        self.assertEqual(parsed_same_row["auction_date"], "2026.09.08")
        self.assertNotEqual(parsed_same_row["auction_date"], "2026.08.01")

        deadline_only = BeautifulSoup(
            """
            <div id="dtl_table">
              <table><tr><th>배당요구종기일</th><td>2026.10.12</td></tr></table>
            </div>
            """,
            "html.parser",
        )
        parsed_deadline_only = crawler.parse_myauction_detail(deadline_only, "https://www.my-auction.co.kr/view/test")
        self.assertEqual(parsed_deadline_only["auction_date"], "")
        self.assertEqual(parsed_deadline_only["입찰기일"], "")

    def test_current_auction_round_supplies_date_min_price_and_deposit_together(self):
        html = """
        <html><body>
          <div id="detail_left">
            <table class="tbl_detail">
              <tr><th>경매종류</th><td>부동산임의경매</td><th>감정가</th><td>847,000,000원</td></tr>
              <tr><th>최저가</th><td>847,000,000원</td><th>입찰보증금</th><td>84,700,000원</td></tr>
              <tr><th>배당요구종기일</th><td>2026.03.26</td></tr>
            </table>
            <h3>기일내역</h3>
            <table>
              <tr><th>회차</th><th>매각기일</th><th>최저매각가격</th><th>저가비율</th><th>매수신청보증금</th><th>결과</th></tr>
              <tr><td>1</td><td>2026.08.04</td><td>847,000,000원</td><td>100%</td><td>84,700,000원</td><td>유찰</td></tr>
              <tr><td>2</td><td>2026.09.08</td><td>592,900,000원</td><td>70%</td><td>59,290,000원</td><td>진행</td></tr>
            </table>
          </div>
          <p class="plan_day"><span class="pink">배당요구종기일 2026.03.26</span></p>
        </body></html>
        """
        parsed = crawler.parse_myauction_detail(BeautifulSoup(html, "html.parser"), "https://www.my-auction.co.kr/view/test")
        self.assertEqual(parsed["auction_date"], "2026.09.08")
        self.assertEqual(parsed["입찰기일"], "2026.09.08")
        self.assertEqual(parsed["min_price"], "592,900,000원")
        self.assertEqual(parsed["deposit"], "59,290,000원")
        self.assertEqual(parsed["min_rate"], "70%")

    def test_current_round_min_price_is_not_overwritten_by_header_summary_amount(self):
        html = """
        <html><body>
          <div id="detail_left">
            <table class="tbl_detail">
              <tr><th>경매종류</th><td>부동산임의경매</td><th>감정가</th><td>1,847,000,000원</td></tr>
              <tr><th>최저가</th><td>3,030,000,000원</td><th>입찰보증금</th><td>303,000,000원</td></tr>
            </table>
            <h3>기일내역</h3>
            <table>
              <tr><th>회차</th><th>매각기일</th><th>최저매각가격</th><th>저가비율</th><th>매수신청보증금</th><th>결과</th></tr>
              <tr><td>1</td><td>2026.09.29</td><td>1,847,000,000원</td><td>100%</td><td>184,700,000원</td><td>진행</td></tr>
            </table>
          </div>
          <div id="dtl_title">
            <h3>사건 정보</h3>
            <div>기타</div>
            <div>
              <ul>
                <li><span class="sale_txt"><span>2026.09.29 10:30</span></span> 최저가 3,030,000,000원 입찰보증금 303,000,000원</li>
              </ul>
            </div>
          </div>
        </body></html>
        """
        parsed = crawler.parse_myauction_detail(BeautifulSoup(html, "html.parser"), "https://www.my-auction.co.kr/view/test")
        self.assertEqual(parsed["min_price"], "1,847,000,000원")
        self.assertEqual(parsed["deposit"], "184,700,000원")
        self.assertEqual(parsed["min_rate"], "100%")

    def test_detail_table_authoritative_cells_override_stale_round_and_header_values(self):
        filler_rows = "\n".join("<tr><td></td><td></td><td></td></tr>" for _ in range(16))
        html = f"""
        <html><body>
          <div id="detail_left">
            <table class="tbl_detail">
              <tr><th>경매종류</th><td>부동산임의경매</td><th>감정가</th><td>206,000,000원</td></tr>
              <tr><th>최저가</th><td>206,000,000원</td><th>입찰보증금</th><td>20,600,000원</td></tr>
            </table>
          </div>
          <p class="plan_day"><span class="pink">매각기일 2026.09.08</span></p>
          <div id="dtl_table">
            <table><tbody>
              <tr><td></td><td></td><td></td></tr>
              <tr><th>최저가</th><td class="tdl_right"><strong>25,998,000원</strong></td><td></td></tr>
              {filler_rows}
              <tr><td>입찰기일</td><td></td><td>2026.10.13</td></tr>
            </tbody></table>
          </div>
        </body></html>
        """
        parsed = crawler.parse_myauction_detail(BeautifulSoup(html, "html.parser"), "https://www.my-auction.co.kr/view/test")
        self.assertEqual(parsed["min_price"], "25,998,000원")
        self.assertEqual(parsed["auction_date"], "2026.10.13")
        self.assertEqual(parsed["입찰기일"], "2026.10.13")
        self.assertNotEqual(parsed["min_price"], "206,000,000원")
        self.assertNotEqual(parsed["auction_date"], "2026.09.08")

    def test_live_detail_dom_overrides_stale_page_source_values(self):
        class LiveDetailDriver:
            def execute_script(self, _script, selector):
                values = {
                    "#dtl_table > table > tbody > tr:nth-child(2) > td.tdl_right > strong": "25,998,000원",
                    "#dtl_table > table > tbody > tr:nth-child(19) > td:nth-child(3)": "2026.10.13",
                    "#dtl_table > table > tbody > tr:nth-child(3) > td.tdl_right": "2,599,800원",
                }
                return values.get(selector, "")

        html = """
        <html><body>
          <div id="detail_left">
            <table class="tbl_detail">
              <tr><th>경매종류</th><td>부동산임의경매</td><th>감정가</th><td>206,000,000원</td></tr>
              <tr><th>최저가</th><td>206,000,000원</td><th>입찰보증금</th><td>20,600,000원</td></tr>
            </table>
          </div>
          <p class="plan_day"><span class="pink">매각기일 2026.09.08</span></p>
        </body></html>
        """
        parsed = crawler.parse_myauction_detail(
            BeautifulSoup(html, "html.parser"),
            "https://www.my-auction.co.kr/view/test",
            driver=LiveDetailDriver(),
        )
        self.assertEqual(parsed["min_price"], "25,998,000원")
        self.assertEqual(parsed["deposit"], "2,599,800원")
        self.assertEqual(parsed["auction_date"], "2026.10.13")
        self.assertEqual(parsed["입찰기일"], "2026.10.13")

    def test_live_sale_header_date_wins_over_stale_fixed_detail_row(self):
        class LiveDetailDriver:
            def execute_script(self, script, selector):
                if "outerHTML" in script and selector == "#dtl_table":
                    return """
                    <div id="dtl_table"><table><tbody>
                      <tr><td></td><td></td><td></td></tr>
                      <tr><th>최저가</th><td class="tdl_right"><strong>9,365,925,000원</strong></td><td></td></tr>
                      <tr><th>입찰보증금</th><td class="tdl_right">936,592,500원</td><td></td></tr>
                      <tr><td>배당요구종기일</td><td></td><td>2025.11.05</td></tr>
                      <tr><td>기타</td><td></td><td>2026.09.22</td></tr>
                    </tbody></table></div>
                    """
                values = {
                    "#dtl_title > div:nth-child(3) > ul > li > span.sale_txt > span": "2026.09.30 10:00",
                    "#dtl_title .sale_txt span": "2026.09.30 10:00",
                    "#dtl_table > table > tbody > tr:nth-child(2) > td.tdl_right > strong": "9,365,925,000원",
                    "#dtl_table > table > tbody > tr:nth-child(19) > td:nth-child(3)": "2026.09.22",
                    "#dtl_table > table > tbody > tr:nth-child(3) > td.tdl_right": "936,592,500원",
                }
                return values.get(selector, "")

        html = """
        <html><body>
          <p class="plan_day"><span class="pink">매각기일 2026.09.22</span></p>
        </body></html>
        """
        parsed = crawler.parse_myauction_detail(
            BeautifulSoup(html, "html.parser"),
            "https://www.my-auction.co.kr/view/test",
            driver=LiveDetailDriver(),
        )
        self.assertEqual(parsed["min_price"], "9,365,925,000원")
        self.assertEqual(parsed["deposit"], "936,592,500원")
        self.assertEqual(parsed["auction_date"], "2026.09.30 10:00")
        self.assertNotEqual(parsed["auction_date"], "2026.09.22")

    def test_sale_spec_pdf_text_keeps_multiple_tenant_rows(self):
        text = """
        점유자 성명 점유부분 정보출처구분 점유의 권원 임대차기간 보증금 차임 전입신고일자 확정일자 배당요구여부 (배당요구일자)
        고정수 현황조사 주거 임차인 2024.06.10.
        김대현 현황조사 주거 임차인 2024.03.15.
        김수형 현황조사 주거 임차인 2024.10.21.
        안의재 현황조사 주거 임차인 2025.04.08.
        이은우 현황조사 주거 임차인 2020.05.18.
        이희석 현황조사 주거 임차인 2024.04.02.
        조영민 2층 301호 등기사항전부증명서 주거 임차인 2020.04.15. 140,000,000 350,000 2020.04.14. 2020.04.16.
        2층 301호 권리신고 주거 임차인 2020.04.15. 140,000,000 350,000 2020.04.14. 2020.04.16. 2025.9.16.
        등기된 부동산에 관한 권리 또는 가처분으로 매각으로 그 효력이 소멸되지 아니하는 것
        """
        tenants = parse_sale_spec_tenants_from_pdf_text(text)
        self.assertEqual([tenant["name"] for tenant in tenants], [
            "고정수",
            "김대현",
            "김수형",
            "안의재",
            "이은우",
            "이희석",
            "조영민",
        ])
        self.assertEqual(tenants[-1]["deposit"], 140_000_000)
        self.assertEqual(tenants[-1]["rent"], 350_000)
        self.assertEqual(tenants[-1]["moveInDate"], "2020.04.14")
        self.assertEqual(tenants[-1]["fixedDate"], "2020.04.16")
        self.assertEqual(tenants[-1]["depositClaimDate"], "2025.09.16")

        opinion_data = briefing_rights.build_opinion_data({
            "rights": [{"type": "근저당권", "date": "2024.03.21", "creditor": "농업협동조합자산관리회사", "amount": "1,000,000원", "isBaseRight": True}],
            "tenants": tenants,
            "tenant_source": "sale_spec_ocr",
            "tenant_ocr_text": text,
            "sale_spec_dividend_deadline": "2025.11.05",
            "min_price": "9,365,925,000원",
            "claim_amount": "2,477,462,227원",
            "auction_type": "부동산임의경매",
        })
        tenant_text = opinion_data["tenantAnalysisText"]
        for name in ("고정수", "김대현", "김수형", "안의재", "이은우", "이희석", "조영민"):
            self.assertIn(name, tenant_text)

    def test_sale_spec_ocr_parse_wins_over_single_pdf_fallback(self):
        collapsed_pdf_text = """
        매각물건명세서
        점유자 성명 점유부분 정보출처구분 점유의 권원 임대차기간 보증금 차임 전입신고일자 확정일자 배당요구여부 (배당요구일자)
        고정수
        현황조사
        주거
        임차인
        2020.04.15.
        140,000,000
        350,000
        2020.04.14.
        2020.04.16.
        비고
        """
        ocr_text = """
        매각물건명세서
        점유자 성명 점유부분 정보출처구분 점유의 권원 임대차기간 보증금 차임 전입신고일자 확정일자 배당요구여부 (배당요구일자)
        고정수 현황조사 주거 임차인 2024.06.10.
        김대현 현황조사 주거 임차인 2024.03.15.
        김수형 현황조사 주거 임차인 2024.10.21.
        안의재 현황조사 주거 임차인 2025.04.08.
        이은우 현황조사 주거 임차인 2020.05.18.
        이희석 현황조사 주거 임차인 2024.04.02.
        조영민 현황조사 주거 임차인 2020.04.14.
        등기된 부동산에 관한 권리 또는 가처분으로 매각으로 그 효력이 소멸되지 아니하는 것
        """

        with (
            patch(
                "app.services.rights_certificate.collect_sale_spec_text_and_images",
                return_value=(collapsed_pdf_text, ["page1.png"]),
            ),
            patch("app.services.rights_certificate.ocr_image_to_text", return_value=ocr_text),
        ):
            context = extract_sale_spec_tenant_context_by_ocr(object(), task_id="multi-tenant-ocr")

        self.assertEqual(
            [tenant["name"] for tenant in context["tenants"]],
            ["고정수", "김대현", "김수형", "안의재", "이은우", "이희석", "조영민"],
        )
        self.assertEqual(context["tenants"][0].get("deposit") or 0, 0)

    def test_no_tenant_sale_spec_ocr_does_not_promote_header_noise(self):
        ocr_text = """
        매각물건명세서
        점유자 성명 점유부분 정보출처구분 점유의 권원 임대차기간 보증금 차임 전입신고일자 확정일자 배당요구여부 (배당요구일자)
        ※ 조사된 임차내역이 없습니다.
        부동산의 점유자와 점유의 권원
        점유자 성명: 와 점유의 AA / 점유구분: 자와 점유의 AA / 보증금: 담당자 확인 필요 / 차임: 없음 또는 미확인
        신고일자
        Be
        ya 기간
        일자-사업자등
        2026.11.04
        """

        self.assertTrue(parse_sale_spec_tenants_from_ocr(ocr_text)[0]["name"].startswith("조사된 임차내역 없음"))
        self.assertTrue(parse_sale_spec_tenants_from_pdf_text(ocr_text)[0]["name"].startswith("조사된 임차내역 없음"))
        self.assertTrue(parse_tenants_from_ocr(ocr_text)[0]["name"].startswith("조사된 임차내역 없음"))

        data = build_template_data({
            "rights": [{"type": "근저당권", "date": "2024.01.10", "creditor": "테스트은행", "isBaseRight": True}],
            "tenants": parse_sale_spec_tenants_from_ocr(ocr_text),
            "tenant_ocr_text": ocr_text,
        })
        self.assertEqual(data["tenantAnalysisText"], "조사된 임차인이 없으므로, 매수인에게 인수되는 임차권리는 없습니다.")
        for noise in ("와 점유의 AA", "신고일자", "Be", "ya", "일자-사업자등", "미확인 점유자"):
            self.assertNotIn(noise, data["tenantAnalysisText"])

    def test_s1_briefing_blocks_false_special_tags_and_internal_markers(self):
        rights = [
            {"type": "근저당권", "date": "2023.05.16", "creditor": "테스트은행", "amount": "500,000,000원", "isBaseRight": True},
            {"type": "강제경매", "date": "2026.01.10", "creditor": "테스트채권자", "amount": "761,315,340원", "isAuctionProcedure": True},
        ]
        data = {
            "auction_type": "부동산강제경매",
            "claim_amount": "761,315,340원",
            "min_price": "592,900,000원",
            "rights": rights,
            "tenants": [{"name": "임차인 없음"}],
            "source_completeness": {"registry": True},
            "sale_spec_remarks": (
                "양식 안내: 가등기담보권, 가압류, 전세권의 등기일자가 말소기준권리보다 빠른 경우 확인. "
                "가처분으로 매각으로 그 효력이 소멸되지 아니하는 것."
            ),
            "status_survey_etc": "특이사항 없음",
            "case_notice": "checklistRows 판정식 YAML review-standard {내부 확인}",
            "management_fee": {},
        }
        opinion_data = briefing_rights.build_opinion_data(data)
        special = briefing_opinion.build_special_opinion(opinion_data)
        self.assertNotIn("가등기", special)
        self.assertNotIn("가처분", special)
        self.assertNotIn("{", special)
        self.assertNotIn("}", special)
        self.assertNotIn("checklistRows", special)
        self.assertNotIn("판정식", special)
        self.assertEqual(registered_right_amount_total(rights), 500_000_000)
        self.assertIn("등기상 권리 기재금액 합계는 500,000,000원", opinion_data["surplusDescription"])

    def test_management_fee_scraper_rejects_polluted_schedule_block(self):
        polluted = """
        <html><body>
          <div id="dtl_stock">
            <div id="dtl_title"><h3>기일내역</h3></div>
            <table>
              <tr><th>저가 비율</th><th>상태</th><th>날짜</th><th>회차</th><th>최저가</th><th>결과</th></tr>
              <tr><td>100%</td><td>284일</td><td>2026-08-04</td><td>1</td><td>847,000,000원</td><td>유찰</td></tr>
            </table>
            <div>미납관리비 목록 (0000.00.00현재) 감정평가현황 목록</div>
          </div>
        </body></html>
        """
        self.assertEqual(_extract_management_fee(BeautifulSoup(polluted, "html.parser")), {})

        clean = """
        <html><body>
          <div id="dtl_stock">
            <div id="dtl_title"><h3>미납관리비</h3></div>
            <table><tr><th>체납관리비</th><td>1,234,000원</td></tr></table>
          </div>
        </body></html>
        """
        extracted = _extract_management_fee(BeautifulSoup(clean, "html.parser"))
        self.assertEqual(extracted["unpaidAmount"], 1_234_000)
        self.assertEqual(extracted["amountStatus"], "confirmed")

    def test_case_1232734_registry_claim_and_ocr_are_canonicalized(self):
        soup = BeautifulSoup(
            """
            <table><tr><th>경매종류</th><td>부동산임의경매</td><th>매각기일</th><td>2026.09.08</td><th>청구금액</th><td>514,682,600원</td></tr></table>
            <div id="dtl_stock">
              <div id="dtl_title"><h3>건물 등기부현황</h3></div>
              <table>
                <tr><th>순위</th><th>등기일</th><th>권리종류</th><th>권리자</th><th>금액</th><th>비고</th></tr>
                <tr><td>1</td><td>2019.01.25</td><td>근저당권</td><td>중소기업은행</td><td>513,000,000원</td><td>말소기준</td></tr>
                <tr><td>2</td><td>2021.10.18</td><td>근저당권</td><td>상상인저축은행</td><td>600,000,000원</td><td>소멸</td></tr>
                <tr><td>3</td><td>2022.06.02</td><td>임의경매</td><td>웰릭스에프앤아이대부</td><td>청구금액 514,682,600원</td><td>2022타경102285 소멸</td></tr>
              </table>
            </div>
            """,
            "html.parser",
        )
        html_rights = _extract_rights(soup)
        self.assertEqual(len(html_rights), 3)
        ocr_rights = parse_rights_from_ocr(
            "2019.01.25 근저당권 중소기업은행 513,000,000원 말소기준\n"
            "2021.10.18 근저당권 상상 600,000,000원 소멸\n"
            "2022.06.02 임의경매 ES 514,682,600원 소멸"
        )
        rights = merge_rights(html_rights, ocr_rights)
        self.assertEqual(len(rights), 3)
        self.assertEqual(len(substantive_registered_rights(rights)), 2)
        self.assertEqual(registered_right_amount_total(rights), 1_113_000_000)

        data = build_template_data({
            "case_number": "2022타경102285",
            "court": "안양지원 경매4계",
            "item_type": "아파트",
            "auction_type": "부동산임의경매",
            "claim_amount": "514,682,600원",
            "min_price": "760,800,000원",
            "rights": rights,
            "tenants": parse_tenants_from_ocr("※ 조사된 임차내역이 없습니다"),
            "tenant_ocr_text": "※ 조사된 임차내역이 없습니다",
            "source_completeness": {"registry": True},
        })
        self.assertEqual(data["court"], "수원지방법원 안양지원")
        self.assertEqual(data["registeredRightCount"], 2)
        self.assertEqual(data["registeredRightAmountTotal"], 1_113_000_000)
        self.assertEqual(data["auctionProcedureCount"], 1)
        self.assertIn("2019.01.25 일자 근저당권 [중소기업은행]", data["baseRightDescription"])
        self.assertIn("매수인에게 인수되는 등기상 권리는 없습니다", data["baseRightDescription"])
        self.assertEqual(data["tenantAnalysisText"], "조사된 임차인이 없으므로, 매수인에게 인수되는 임차권리는 없습니다.")
        self.assertIn("1,113,000,000원", data["narrativeReportHtml"])
        self.assertIn("청구금액은 514,682,600원", data["narrativeReportHtml"])
        self.assertIn("법원 경매절차비용을 차감하기 전", data["surplusDescription"])
        self.assertIn("가장 후순위 담보권을 신청담보권으로 보는 보수적 기준", data["surplusDescription"])
        self.assertIn("선순위인 기재금액은 513,000,000원", data["surplusDescription"])
        self.assertIn("배당재원은 247,800,000원", data["surplusDescription"])
        self.assertNotIn("선순위인 기재금액은 담당자 확인 필요", data["surplusDescription"])
        self.assertIn("무잉여 가능성은 사실상 없습니다", data["surplusDescription"])
        self.assertIn("신청담보권에도 1원 이상 배당될 수 있어", data["surplusDescription"])
        self.assertNotIn("현재 기재값 기준", data["surplusDescription"])
        self.assertIn("법원 기록으로 최종 확인해야 합니다", data["surplusDescription"])
        self.assertNotIn("전체 수집 완결성", data["narrativeReportHtml"])
        self.assertNotIn("같은 날짜에 설정된 권리는", data["narrativeReportHtml"])

    def test_forced_auction_entry_can_be_base_right_without_senior_takeover_caveats(self):
        soup = BeautifulSoup(
            """
            <div id="dtl_stock">
              <div id="dtl_title"><h3>건물 등기부현황</h3></div>
              <table>
                <tr><th>순위</th><th>등기일</th><th>권리종류</th><th>권리자</th><th>금액</th><th>비고</th></tr>
                <tr><td>1</td><td>2025.05.09</td><td>강제경매</td><td></td><td>90,000,000원</td><td>말소기준</td></tr>
              </table>
            </div>
            """,
            "html.parser",
        )
        context = extract_rights_context(soup)
        rights = context["rights"]
        self.assertTrue(context["source_completeness"]["registry"])
        self.assertEqual(len(rights), 1)
        self.assertEqual(rights[0]["type"], "강제경매")
        self.assertEqual(rights[0]["creditor"], "")
        self.assertEqual(len(substantive_registered_rights(rights)), 0)

        data = build_template_data({
            "case_number": "2025타경101054",
            "auction_type": "부동산강제경매",
            "min_price": "300,000,000원",
            **context,
            "expected_dividend": {
                "auctionApplicantDividendFound": True,
                "auctionApplicantDividendAmount": 1,
                "source": "expected_dividend_table",
            },
            "case_notice": "가등기 없음 가처분 없음",
            "tenant_ocr_text": "조사된 임차내역이 없습니다",
        })

        self.assertIn("2025.05.09 일자 강제경매 기입등기", data["baseRightDescription"])
        self.assertNotIn("[담당자 확인 필요]", data["baseRightDescription"])
        self.assertNotIn("등기 원본 전체 확보 여부", data["baseRightDescription"])
        self.assertIn("매수인에게 인수되는 등기상 권리는 없습니다", data["baseRightDescription"])
        self.assertIn("등기상 선순위 인수 권리는 없습니다", data["narrativeReportHtml"])
        self.assertIn("예상배당표상 말소기준권리 또는 경매신청채권자에게 채권배당금 1원", data["surplusDescription"])
        self.assertIn("현재 예상배당표 기준 무잉여 가능성은 없습니다", data["surplusDescription"])
        self.assertNotIn("입력자료 기재값", data["surplusDescription"])
        self.assertNotIn("가등기(담보/순위보전)</div>", data["narrativeReportHtml"])
        self.assertNotIn("가처분</div>", data["narrativeReportHtml"])

    def test_voluntary_auction_fallback_reads_like_a_written_opinion(self):
        rights = [{
            "type": "근저당권",
            "date": "2025.04.10",
            "creditor": "신한은행",
            "amount": 72_028_201,
            "isBaseRight": True,
        }]
        text = analyze_surplus(
            {
                "auction_type": "부동산임의경매",
                "claim_amount": "82,859,092원",
                "min_price": "1,376,800,000원",
                "auction_applicant_creditors": ["신한은행"],
            },
            rights,
            rights[0],
            [],
            rights_source_complete=False,
        )
        self.assertIn("본 경매는 근저당권자인 신한은행(청구금액 82,859,092원)이 임의경매를 신청한 사건입니다", text)
        self.assertIn("신청채권자에게 배당될 가능성이 있어 무잉여 가능성은 낮게 판단됩니다", text)
        self.assertIn("취하 가능성은 낮은 편으로 판단됩니다", text)
        self.assertNotIn("입력자료 기재값", text)
        self.assertNotIn("무잉여 산식을 확정하지 않았습니다", text)
        self.assertNotIn("단순 비율로 채권자의 향후 취하 여부를 예측하지 않습니다", text)

    def test_myungseung_analysis_is_scoped_and_rendered_as_a_special_issue(self):
        soup = BeautifulSoup(
            """
            <div id="dtl_stock"><div id="dtl_title"><h3>관련사건</h3></div><div class="excmt"><table class="tbl_excmt"><tr><th>재진행</th><td>무관한 문구</td></tr></table></div></div>
            <div id="dtl_stock"><div id="dtl_title"><h3>법무법인 명승 권리분석</h3></div><div class="excmt"><table class="tbl_excmt"><tr><th>재진행</th><td>중단되었던 매각 절차가 재개되었으며 입찰 전 최신 등기와 점유를 재확인해야 합니다.</td></tr></table></div></div>
            """,
            "html.parser",
        )
        analysis = _extract_myungseung_rights_analysis(soup)
        self.assertEqual(len(analysis), 1)
        self.assertEqual(analysis[0]["label"], "재진행")
        self.assertNotIn("무관한 문구", analysis[0]["text"])
        data = build_template_data({
            "case_number": "2022타경102285",
            "item_type": "아파트",
            "rights": [],
            "tenants": [],
            "myungseung_analysis": analysis,
        })
        report = data["narrativeReportHtml"]
        self.assertIn("<strong>재진행</strong>", report)
        self.assertIn("마이옥션 상세페이지 내 법무법인 명승 권리분석", report)
        self.assertNotIn("법무법인 명승 권리분석 · 법무법인 명승 권리분석", report)
        self.assertIn("중단되었던 매각 절차가 재개", report)

    def test_powerpoint_pdf_export_initializes_com_in_worker_thread(self):
        app = MagicMock()
        presentation = MagicMock()
        app.Presentations.Open.return_value = presentation

        with tempfile.TemporaryDirectory() as temp_dir:
            pptx_path = Path(temp_dir) / "certificate.pptx"
            pdf_path = Path(temp_dir) / "certificate.pdf"
            pptx_path.write_bytes(b"pptx")
            presentation.SaveAs.side_effect = lambda path, _format: Path(path).write_bytes(b"%PDF-1.7")

            with (
                patch("pythoncom.CoInitialize") as initialize,
                patch("pythoncom.CoUninitialize") as uninitialize,
                patch("win32com.client.DispatchEx", return_value=app) as dispatch,
            ):
                self.assertTrue(export_pptx_to_pdf(pptx_path, pdf_path))

        initialize.assert_called_once_with()
        dispatch.assert_called_once_with("PowerPoint.Application")
        presentation.Close.assert_called_once_with()
        app.Quit.assert_called_once_with()
        uninitialize.assert_called_once_with()

    def test_first_page_real_case_info_stays_two_lines_inside_fixed_box(self):
        template_path = BACKEND_PATH / "templates" / "rights_certificate" / "certificate.pptx"
        source = Presentation(str(template_path))
        source_case_shape = next(
            shape for shape in source.slides[0].shapes
            if getattr(shape, "has_text_frame", False) and "{{caseNumber}}" in shape.text
        )
        first_section_top = min(
            shape.top for shape in source.slides[0].shapes
            if getattr(shape, "has_table", False) and shape.top > source_case_shape.top
        )
        data = build_template_data(
            {
                "case_number": "2026타경12345",
                "court": "서울중앙지방법원",
                "item_type": "아파트",
                "rights": [],
                "tenants": [],
            }
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "case-info-fit.pptx"
            render_certificate_pptx_template(template_path, output_path, data)
            rendered = Presentation(str(output_path))

        case_shape = next(
            shape for shape in rendered.slides[0].shapes
            if getattr(shape, "has_text_frame", False) and "2026타경12345" in shape.text
        )
        self.assertEqual(case_shape.top, source_case_shape.top)
        self.assertEqual(case_shape.height, source_case_shape.height)
        self.assertLess(case_shape.top + case_shape.height, first_section_top)
        self.assertEqual(
            [paragraph.text for paragraph in case_shape.text_frame.paragraphs],
            ["사건번호: 2026타경12345", "서울중앙지방법원"],
        )
        self.assertFalse(case_shape.text_frame.word_wrap)
        self.assertEqual(case_shape.text_frame.auto_size, MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE)
        self.assertLessEqual(
            max(
                run.font.size.pt
                for paragraph in case_shape.text_frame.paragraphs
                for run in paragraph.runs
            ),
            12,
        )

    def test_pptx_keeps_first_page_and_replaces_tables_with_clean_narrative(self):
        template_path = BACKEND_PATH / "templates" / "rights_certificate" / "certificate.pptx"
        source = Presentation(str(template_path))
        source_first_shape_count = len(source.slides[0].shapes)
        data = build_template_data(
            {
                "case_number": "2026타경12345",
                "court": "서울중앙지방법원",
                "item_type": "아파트",
                "rights": [
                    {"type": "근저당권", "date": "2024-01-10", "creditor": "테스트은행", "isBaseRight": True},
                    {"type": "압류", "date": "2024-02-10", "creditor": "서울시"},
                ],
                "tenants": [{"name": "임차인 없음"}],
                "tenant_ocr_text": "임차인 없음",
                "sale_spec_remarks": "특별매각조건 해당 사항 없음",
                "case_document_text": "문건처리내역 확인 완료",
                "management_fee": {
                    "unpaidAmount": 0,
                    "amountStatus": "none",
                    "note": "미납 관리비 없음",
                },
                "source_completeness": {
                    "registry": True,
                    "sale_spec": True,
                    "status_survey": True,
                    "case_documents": True,
                    "appraisal": True,
                    "building_register": True,
                    "dividend_requests": True,
                },
                "author_name": "테스트 담당자",
            }
        )

        self.assertGreaterEqual(len(data["narrativePages"]), 2)
        self.assertNotIn("checklistRows", data)
        self.assertNotIn("checklistDetails", data)
        self.assertNotIn("specialSummaryText", data)

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "certificate-with-narrative.pptx"
            render_certificate_pptx_template(template_path, output_path, data)
            rendered = Presentation(str(output_path))

        self.assertEqual(len(rendered.slides), 2 + len(data["narrativePages"]))
        self.assertEqual(len(rendered.slides[0].shapes), source_first_shape_count)
        self.assertIn("1. 말소기준 및 등기부상 소멸사항", _slide_text(rendered.slides[0]))
        narrative_text = "\n".join(_slide_text(slide) for slide in list(rendered.slides)[1:-1])
        self.assertIn("4. 특이사항", narrative_text)
        self.assertIn("1) 종합 판단", narrative_text)
        self.assertNotIn("특이사항 종합 검토의견", narrative_text)
        self.assertNotIn("종합 결론 및 잔여 확인사항", narrative_text)
        self.assertGreater(len(narrative_text), 800)
        self.assertNotIn("특이사항 종합 검토의견", narrative_text)
        self.assertNotIn("자료 미확보 및 추가 확인 범위", narrative_text)
        prototype_hashes = _slide_image_hashes(source.slides[1])
        closing_hashes = _slide_image_hashes(source.slides[-1])
        self.assertTrue(prototype_hashes)
        self.assertTrue(closing_hashes)
        self.assertEqual(_slide_image_hashes(rendered.slides[-1]), closing_hashes)
        for slide in list(rendered.slides)[1:-1]:
            self.assertEqual(_slide_image_hashes(slide), prototype_hashes)
            tables = [shape for shape in slide.shapes if getattr(shape, "has_table", False)]
            self.assertEqual(len(tables), 2)
        for forbidden in (
            "특이사항 종합 체크표", "YAML", "review-standard", "규칙엔진", "이상없음 17",
            "관리비 채권의 소멸시효는 3년",
        ):
            self.assertNotIn(forbidden, narrative_text)
        self.assertNotIn("{{", "\n".join(_slide_text(slide) for slide in rendered.slides))

    def test_pptx_narrative_refuses_landscape_template_before_writing_output(self):
        data = build_template_data(
            {
                "case_number": "2026타경54321",
                "item_type": "아파트",
                "rights": [],
                "tenants": [],
                "management_fee": {"unpaidAmount": 0},
            }
        )
        landscape = Presentation()
        landscape.slide_width = Inches(13.333)
        landscape.slide_height = Inches(7.5)
        landscape.slides.add_slide(landscape.slide_layouts[6])

        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "landscape-template.pptx"
            output_path = Path(temp_dir) / "must-not-be-written.pptx"
            landscape.save(str(template_path))

            with self.assertRaisesRegex(
                ValueError,
                r"최소 7\.5 x 10\.8인치.*Narrative report requires slides at least 7\.5 x 10\.8 inches",
            ):
                render_certificate_pptx_template(template_path, output_path, data)

            self.assertFalse(output_path.exists())

    def test_pptx_expands_long_risky_report_without_blank_or_table_slides(self):
        template_path = BACKEND_PATH / "templates" / "rights_certificate" / "certificate.pptx"
        data = build_template_data(
            {
                "case_number": "2026타경59999",
                "item_type": "토지",
                "rights": [{"type": "근저당권", "date": "2024-01-10", "isBaseRight": True}],
                "tenants": [{"name": "임차인 없음"}],
                "case_document_text": (
                    "신탁등기 대위변제 분묘기지권 유치권 법정지상권 위반건축물 "
                    "농지취득자격증명 지분매각 토지 별도등기 대지권 미등기"
                ),
            }
        )
        self.assertGreater(len(data["narrativePages"]), 3)

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "long-risk-report.pptx"
            render_certificate_pptx_template(template_path, output_path, data)
            rendered = Presentation(str(output_path))

        self.assertEqual(len(rendered.slides), 2 + len(data["narrativePages"]))
        for slide in list(rendered.slides)[1:-1]:
            self.assertTrue(_slide_text(slide).strip())
            self.assertEqual(
                sum(bool(getattr(shape, "has_table", False)) for shape in slide.shapes),
                2,
            )

    def test_risky_narrative_has_only_issue_blocks_and_no_internal_sources(self):
        data = build_template_data(
            {
                "case_number": "2026타경67890",
                "court": "서울중앙지방법원",
                "item_type": "아파트",
                "rights": [
                    {"type": "근저당권", "date": "2024-01-10", "creditor": "A은행", "isBaseRight": True},
                ],
                "tenants": [{
                    "name": "홍길동",
                    "moveInDate": "2023-01-01",
                    "fixedDate": "",
                    "depositClaimDate": "",
                    "deposit": "100,000,000원",
                    "occupancyType": "주거",
                }],
                "tenant_ocr_text": "홍길동 전입 2023년 1월 1일 보증금 100,000,000원",
                "case_document_text": "유치권 신고서 접수 및 신탁등기 관련 문건",
                "sale_spec_remarks": "유치권 신고 있음",
                "management_fee": {"unpaidAmount": 2_300_000, "amountStatus": "confirmed"},
                "source_completeness": {
                    "registry": True,
                    "sale_spec": True,
                    "case_documents": True,
                },
            }
        )

        issue_blocks = [
            block
            for page in data["narrativePages"]
            for block in page["blocks"]
            if block.get("sectionTitle") in ("권리관계 상세 검토", "임차·점유 및 보증금 인수 검토")
        ]
        issue_headings = [block["heading"] for block in issue_blocks]
        self.assertIn("유치권", issue_headings)
        self.assertIn("신탁등기", issue_headings)
        self.assertIn("보증금 인수", issue_headings)
        self.assertNotIn("선순위 전세권", issue_headings)
        fee_pages = [
            page for page in data["narrativePages"]
            if any(block["heading"] == "체납관리비" for block in page["blocks"])
        ]
        self.assertEqual(len(fee_pages), 1)
        self.assertGreater(len(fee_pages[0]["blocks"]), 1)
        for block in issue_blocks:
            with self.subTest(issue=block["heading"]):
                self.assertIn("확인", block["body"])
                self.assertNotIn("확인 사실:", block["body"])
                self.assertNotIn("권리 판단:", block["body"])
                self.assertNotIn("예상 영향:", block["body"])
                self.assertNotIn("입찰 전 조치:", block["body"])
        customer_text = data["narrativeReportHtml"]
        self.assertIn("유치권", customer_text)
        for forbidden in (
            "YAML", "review-standard", "권리분석_규칙.json", "32개", "특이사항 종합 체크표",
            "자동 판정", "자동 확인", "자동 탐지", "문건 탐지",
        ):
            self.assertNotIn(forbidden, customer_text)

    def test_missing_sources_never_become_clean_or_no_takeover(self):
        data = build_template_data(
            {"case_number": "2026타경1", "item_type": "아파트", "rights": [], "tenants": []}
        )
        text = data["narrativeReportHtml"]
        self.assertIn("등기부현황 확인이 필요한 사건입니다", text)
        self.assertIn("임차·점유 확인이 필요한 사건입니다", text)
        self.assertNotIn("깨끗하다고 단정하지 않습니다", text)
        self.assertNotIn("인수되는 임차권리는 없습니다", text)
        self.assertNotIn("인수되는 임차권리는 없습니다", data["tenantAnalysisText"])
        self.assertNotIn("인수 없음", data["tenantAnalysisText"])

        items = build_checklist_from_pipeline(
            data={}, rights=[], base_right=None, valid_tenants=[], management_fee={}, texts=[]
        )
        by_name = {item.name: item for item in items}
        self.assertEqual(by_name["선순위 전세권"].state, State.UNKNOWN)
        self.assertEqual(by_name["대항력"].state, State.UNKNOWN)
        self.assertEqual(by_name["보증금 인수"].state, State.UNKNOWN)
        self.assertEqual(by_name["체납관리비"].state, State.UNKNOWN)

    def test_explicit_no_tenant_row_is_distinct_from_empty_input(self):
        explicit = build_template_data(
            {
                "case_number": "2026타경2",
                "item_type": "아파트",
                "rights": [{"type": "근저당권", "date": "2024-01-10", "isBaseRight": True}],
                "tenants": [{"name": "임차인 없음"}],
            }
        )
        missing = build_template_data(
            {
                "case_number": "2026타경3",
                "item_type": "아파트",
                "rights": [{"type": "근저당권", "date": "2024-01-10", "isBaseRight": True}],
                "tenants": [],
            }
        )
        self.assertFalse(explicit["noTenants"])
        self.assertIn("조사된 임차인은 없으므로", explicit["tenantAnalysisText"])
        self.assertIn("인수되는 임차권리는 없습니다", explicit["tenantAnalysisText"])
        self.assertIn("임차·점유 확인이 필요한 사건입니다", missing["narrativeReportHtml"])

    def test_partial_no_tenant_ocr_is_scope_limited_on_first_page(self):
        partial = build_template_data(
            {
                "case_number": "2026타경20",
                "item_type": "아파트",
                "rights": [],
                "tenants": [],
                "tenant_source": "sale_spec_ocr",
                "tenant_ocr_text": "임차인 없음",
            }
        )
        self.assertFalse(partial["noTenants"])
        self.assertIn("조사된 임차인은 없으므로", partial["tenantAnalysisText"])
        self.assertIn("인수되는 임차권리는 없습니다", partial["tenantAnalysisText"])
        self.assertIn("조사된 임차인은 없으므로", partial["narrativeReportHtml"])

        complete = build_template_data(
            {
                "case_number": "2026타경20",
                "item_type": "아파트",
                "rights": [],
                "tenants": [],
                "tenant_source": "sale_spec_ocr",
                "tenant_ocr_text": "임차인 없음",
                "source_completeness": {"sale_spec": True},
            }
        )
        self.assertTrue(complete["noTenants"])
        self.assertIn("매수인에게 인수되는 임차권리는 없습니다", complete["tenantAnalysisText"])

    def test_no_tenant_ocr_requires_tenant_label_and_adjacent_absence_value(self):
        self.assertTrue(_tenant_ocr_text_indicates_no_tenants("임차인: 없음"))
        self.assertTrue(_tenant_ocr_text_indicates_no_tenants("점유자 성명\n해당사항 없음"))
        self.assertFalse(_tenant_ocr_text_indicates_no_tenants("임차인 홍길동\n확정일자 해당없음"))
        self.assertFalse(_tenant_ocr_text_indicates_no_tenants("홍길동 / 전입일 2023-01-01 / 해당사항 없음"))

        data = build_template_data(
            {
                "case_number": "2026타경21",
                "item_type": "아파트",
                "rights": [],
                "tenants": [],
                "tenant_ocr_text": "임차인 홍길동\n확정일자 해당없음",
            }
        )
        self.assertFalse(data["noTenants"])
        self.assertNotIn("인수되는 임차권리는 없습니다", data["tenantAnalysisText"])
        self.assertIn("임차·점유 확인이 필요한 사건입니다", data["narrativeReportHtml"])

    def test_site_survey_explicit_no_tenant_is_definitive_without_fake_name(self):
        text = "※ 조사된 임차내역이 없습니다\n별지 전입세대열람내역상 채무자(소유자) 세대 전입"
        tenants = parse_tenants_from_ocr(text)
        self.assertEqual(len(tenants), 1)
        self.assertNotEqual(tenants[0]["name"], "조사된")
        data = build_template_data({
            "case_number": "2022타경102285",
            "item_type": "아파트",
            "rights": [{"type": "근저당권", "date": "2019.01.25", "isBaseRight": True}],
            "tenants": tenants,
            "tenant_ocr_text": text,
        })
        self.assertEqual(data["tenantAnalysisText"], "조사된 임차인이 없으므로, 매수인에게 인수되는 임차권리는 없습니다.")
        self.assertIn("조사된 임차인은 없으므로", data["narrativeReportHtml"])
        self.assertIn("인수되는 임차권리는 없습니다", data["narrativeReportHtml"])

    def test_source_label_or_missing_base_date_never_proves_no_tenant_takeover(self):
        source_label_only = build_template_data(
            {
                "case_number": "2026타경31",
                "item_type": "아파트",
                "rights": [],
                "tenants": [],
                "tenant_source": "sale_spec_ocr",
                "tenant_ocr_text": "",
            }
        )
        self.assertIn("담당자 확인이 필요", source_label_only["tenantAnalysisText"])
        self.assertNotIn("인수되는 임차권리는 없습니다", source_label_only["tenantAnalysisText"])
        self.assertIn("임차·점유 확인이 필요한 사건입니다", source_label_only["narrativeReportHtml"])

        missing_base = build_template_data(
            {
                "case_number": "2026타경32",
                "item_type": "아파트",
                "rights": [],
                "tenants": [{
                    "name": "홍길동",
                    "moveInDate": "2023-01-01",
                    "deposit": "100,000,000원",
                    "occupancyType": "주거",
                }],
                "tenant_ocr_text": "홍길동 전입 2023년 1월 1일 보증금 100,000,000원",
            }
        )
        self.assertIn("말소기준권리 일자가 확인되지 않아", missing_base["tenantAnalysisText"])
        self.assertNotIn("인수되는 임차권리는 없습니다", missing_base["tenantAnalysisText"])

    def test_partial_tenant_source_describes_only_the_known_junior_tenant(self):
        data = build_template_data(
            {
                "case_number": "2026타경32",
                "item_type": "아파트",
                "rights": [{"type": "근저당권", "date": "2024-01-10", "isBaseRight": True}],
                "tenants": [{
                    "name": "홍길동",
                    "moveInDate": "2024-02-01",
                    "fixedDate": "2024-02-02",
                    "depositClaimDate": "2024-02-03",
                    "depositDeadline": "2024-03-01",
                }],
            }
        )
        first_page = data["tenantAnalysisText"]
        self.assertIn("대항력이 없으므로(대항력 X)", first_page)
        self.assertIn("낙찰자에게 인수되는 임차권리는 없습니다", first_page)

    def test_missing_tenant_dates_do_not_calculate_priority_repayment_or_safe_dividend(self):
        base_right = {"type": "근저당권", "date": "2024-01-10", "isBaseRight": True}
        data = build_template_data(
            {
                "case_number": "2026타경33",
                "item_type": "아파트",
                "address": "",
                "rights": [base_right],
                "tenants": [{
                    "name": "홍길동",
                    "moveInDate": "",
                    "fixedDate": "",
                    "depositClaimDate": "",
                    "deposit": "50,000,000원",
                    "occupancyType": "주거",
                }],
                "source_completeness": {"registry": True, "sale_spec": True},
            }
        )
        first_page = data["tenantAnalysisText"]
        self.assertIn("최우선변제 여부", first_page)
        self.assertIn("판단을 유보", first_page)
        self.assertNotIn("최우선 변제됩니다", first_page)
        self.assertNotIn("인수되지 않습니다", first_page)

        items = build_checklist_from_pipeline(
            data={"item_type": "아파트"},
            rights=[base_right],
            base_right=base_right,
            valid_tenants=[{"name": "홍길동", "moveInDate": ""}],
            management_fee={},
            tenant_source_confirmed=True,
            rights_source_confirmed=True,
        )
        by_name = {item.name: item for item in items}
        self.assertEqual(by_name["배당요구"].state, State.UNKNOWN)
        self.assertIn("배당요구", data["narrativeReportHtml"])

    def test_late_tenant_dividend_demand_is_risk_and_blocks_priority_amount(self):
        base_right = {"type": "근저당권", "date": "2024-01-10", "isBaseRight": True}
        tenant = {
            "name": "홍길동",
            "moveInDate": "2023-01-01",
            "fixedDate": "2023-01-02",
            "depositClaimDate": "2024-04-01",
            "depositDeadline": "2024-03-01",
            "deposit": "50,000,000원",
            "occupancyType": "주거",
        }
        data = build_template_data(
            {
                "case_number": "2026타경34",
                "item_type": "아파트",
                "address": "서울특별시 중구 세종대로 1",
                "rights": [base_right],
                "tenants": [tenant],
                "source_completeness": {"registry": True, "sale_spec": True},
            }
        )
        first_page = data["tenantAnalysisText"]
        self.assertIn("배당요구일이 배당요구종기보다 늦어", first_page)
        self.assertNotIn("한도에서 우선변제", first_page)
        self.assertNotIn("잔존 보증금은 0원", first_page)

        items = build_checklist_from_pipeline(
            data={"item_type": "아파트"},
            rights=[base_right],
            base_right=base_right,
            valid_tenants=[tenant],
            management_fee={},
            tenant_source_confirmed=True,
            rights_source_confirmed=True,
        )
        by_name = {item.name: item for item in items}
        self.assertEqual(by_name["배당요구"].state, State.RISK)
        self.assertNotIn("적법한 배당요구 제출", by_name["배당요구"].basis)

        separate_request = build_template_data(
            {
                "case_number": "2026타경34-1",
                "item_type": "아파트",
                "address": "서울특별시 중구 세종대로 1",
                "rights": [base_right],
                "tenants": [{**tenant, "depositClaimDate": "", "depositDeadline": ""}],
                "dividend_requests": [{
                    "creditor": "홍길동",
                    "requestDate": "2024-04-01",
                    "deadline": "2024-03-01",
                }],
                "source_completeness": {"registry": True, "sale_spec": True},
            }
        )
        self.assertIn("배당요구일이 배당요구종기보다 늦어", separate_request["tenantAnalysisText"])
        demand_blocks = [
            block for page in separate_request["narrativePages"] for block in page["blocks"]
            if block["heading"] == "배당요구"
        ]
        self.assertTrue(demand_blocks)
        self.assertEqual(demand_blocks[0]["kind"], "risk")

    def test_same_date_right_is_not_described_as_junior(self):
        rights = [
            {"type": "근저당권", "date": "2024-01-10", "creditor": "A은행", "isBaseRight": True},
            {"type": "가처분", "date": "2024-01-10", "creditor": "B"},
        ]
        items = build_checklist_from_pipeline(
            data={}, rights=rights, base_right=rights[0], valid_tenants=[], management_fee={},
            texts=["가처분 등기"], tenant_source_confirmed=False, signal_source_confirmed=True,
        )
        by_name = {item.name: item for item in items}
        self.assertEqual(by_name["가처분"].state, State.CHECK)

        data = build_template_data(
            {"case_number": "2026타경4", "item_type": "아파트", "rights": rights, "tenants": []}
        )
        first_page = data["baseRightDescription"]
        self.assertIn("같은 날짜", first_page)
        self.assertIn("접수번호·순위 확인 전", first_page)
        self.assertNotIn("인수되는 권리는 없습니다", first_page)

    def test_special_situations_outside_internal_32_items_are_not_dropped(self):
        data = build_template_data(
            {
                "case_number": "2026타경41",
                "item_type": "토지",
                "rights": [{"type": "근저당권", "date": "2024-01-10", "isBaseRight": True}],
                "tenants": [{"name": "임차인 없음"}],
                "case_document_text": (
                    "신탁등기가 현재 존속 중. 임차인의 대위변제 가능성이 있음. "
                    "현장에서 분묘 소재가 확인됨."
                ),
            }
        )
        blocks = [block for page in data["narrativePages"] for block in page["blocks"]]
        by_heading = {block["heading"]: block for block in blocks}
        for heading in ("신탁등기", "대위변제 위험", "분묘기지권"):
            with self.subTest(heading=heading):
                self.assertIn(heading, by_heading)
                body = by_heading[heading]["body"]
                self.assertNotIn("확인 사실:", body)
                self.assertNotIn("권리 판단:", body)
                self.assertNotIn("예상 영향:", body)
                self.assertNotIn("입찰 전 조치:", body)

    def test_appraisal_text_is_positive_evidence_but_not_global_absence_coverage(self):
        partial = build_template_data(
            {
                "case_number": "2026타경42",
                "item_type": "아파트",
                "rights": [{"type": "근저당권", "date": "2024-01-10", "isBaseRight": True}],
                "tenants": [],
                "appraisal_raw": "감정평가액 300,000,000원",
            }
        )
        partial_text = partial["narrativeReportHtml"]
        self.assertNotIn("유치권·토지 별도등기 관련 특이 기재는 발견되지 않았습니다", partial_text)
        self.assertNotIn("관련 특이 기재가 발견되지 않았습니다", partial_text)
        self.assertNotIn("자료 미확보 및 추가 확인 범위", partial_text)
        self.assertNotIn("인수되는 권리는 없습니다", partial["baseRightDescription"])

        positive = build_template_data(
            {
                "case_number": "2026타경43",
                "item_type": "아파트",
                "rights": [],
                "tenants": [],
                "appraisal_raw": "감정평가서에 제시외 건물 표시가 있음",
            }
        )
        headings = [block["heading"] for page in positive["narrativePages"] for block in page["blocks"]]
        self.assertIn("제시외 건물", headings)

        timed_out = build_template_data(
            {
                "case_number": "2026타경44",
                "item_type": "아파트",
                "rights": [{"type": "근저당권", "date": "2024-01-10", "isBaseRight": True}],
                "tenants": [],
                "rights_ocr_text": "근저당권 일부 OCR",
                "source_completeness": {"registry": True, "sale_spec": True},
                "source_collection_incomplete": True,
            }
        )
        self.assertNotIn("인수되는 권리는 없습니다", timed_out["baseRightDescription"])
        self.assertNotIn("자료 미확보 및 추가 확인 범위", timed_out["narrativeReportHtml"])

        failed_registry = build_template_data(
            {
                "case_number": "2026타경4-1",
                "item_type": "아파트",
                "rights": [{"type": "근저당권", "date": "2024-01-10", "isBaseRight": True}],
                "source_completeness": {"registry": True},
                "source_status": {
                    "registry": {"collected": True, "complete": True, "failed": True},
                },
            }
        )
        self.assertNotIn("인수되는 권리는 없습니다", failed_registry["baseRightDescription"])

    def test_partial_sources_still_produce_detailed_fact_narrative(self):
        data = build_template_data(
            {
                "case_number": "2026타경35",
                "item_type": "아파트",
                "rights": [
                    {"type": "근저당권", "date": "2024-01-10", "isBaseRight": True},
                    {"type": "압류", "date": "2024-02-10"},
                ],
                "tenants": [{"name": "홍길동", "moveInDate": "2024-03-01"}],
                "appraisal_raw": "감정평가현황 일부 기재",
            }
        )
        report = data["narrativeReportHtml"]
        self.assertIn("실체 권리는 2건", report)
        self.assertIn("임차 관련 기재 1건", report)
        self.assertNotIn("인수되는 권리는 없습니다", data["baseRightDescription"])
        self.assertNotIn("관련 특이 기재는 발견되지 않았습니다", report)

    def test_partial_amount_sources_do_not_produce_surplus_or_withdrawal_conclusions(self):
        data = build_template_data(
            {
                "case_number": "2026타경36",
                "item_type": "아파트",
                "appraised_price": "1,000,000,000원",
                "min_price": "700,000,000원",
                "rights": [{
                    "type": "근저당권",
                    "date": "2024-01-10",
                    "amount": 10_000_000,
                    "isBaseRight": True,
                }],
                "tenants": [],
                "appraisal_raw": "감정가 1,000,000,000원",
            }
        )
        text = data["surplusDescription"]
        self.assertIn("무잉여 여부를 결론내리지 않았습니다", text)
        self.assertIn("단순 비율로 채권자의 향후 취하 여부를 예측하지 않습니다", text)
        self.assertNotIn("입력자료 기재값", text)
        self.assertNotIn("무잉여 가능성은 낮습니다", text)
        self.assertNotIn("감정가 대비 1%", text)

    def test_eviction_execution_cost_does_not_complete_no_surplus_formula(self):
        data = build_template_data(
            {
                "case_number": "2026타경37",
                "item_type": "아파트",
                "appraised_price": "1,000,000,000원",
                "min_price": "700,000,000원",
                "rights": [{
                    "type": "근저당권",
                    "date": "2024-01-10",
                    "amount": 10_000_000,
                    "isBaseRight": True,
                }],
                "tenants": [{"name": "임차인 없음"}],
                "execution_cost_source_complete": True,
                "myauction_eviction_costs": {"grandTotal": 50_000_000},
                "source_completeness": {"registry": True, "appraisal": True},
            }
        )
        text = data["surplusDescription"]
        self.assertIn("무잉여 여부를 결론내리지 않았습니다", text)
        self.assertIn("법원 경매절차비용", text)
        self.assertNotIn("무잉여 가능성은 낮습니다", text)
        self.assertIn("단순 비율로 채권자의 향후 취하 여부를 예측하지 않습니다", text)
        self.assertNotIn("70%", text)

    def test_no_surplus_basis_requires_explicit_court_cost_and_applicant_priority(self):
        rights = [{
            "type": "근저당권",
            "date": "2024-01-10",
            "creditor": "A은행",
            "amount": 300_000_000,
            "isBaseRight": True,
        }]
        old_eviction_basis = calculate_surplus_basis(
            {
                "min_price": "700,000,000원",
                "auction_applicant_creditors": ["A은행"],
                "execution_cost_source_complete": True,
                "myauction_eviction_costs": {"grandTotal": 50_000_000},
            },
            rights,
            rights[0],
        )
        self.assertFalse(old_eviction_basis["can_calculate"])
        self.assertFalse(old_eviction_basis["court_cost_available"])
        self.assertEqual(old_eviction_basis["court_auction_cost"], 0)

        explicit_court_basis = calculate_surplus_basis(
            {
                "min_price": "700,000,000원",
                "auction_applicant_creditors": ["A은행"],
                "court_auction_cost": "10,000,000원",
                "court_auction_cost_source_complete": True,
            },
            rights,
            rights[0],
        )
        self.assertTrue(explicit_court_basis["can_calculate"])
        self.assertEqual(explicit_court_basis["remainder"], 690_000_000)

    def test_disclosed_voluntary_basis_formats_no_prior_rights_as_zero_won(self):
        rights = [
            {
                "type": "근저당권",
                "date": "2024-01-10",
                "creditor": "A은행",
                "amount": 300_000_000,
                "isBaseRight": True,
            },
            {
                "type": "임의경매",
                "date": "2024-06-01",
                "creditor": "A은행",
                "amount": 301_000_000,
                "isAuctionProcedure": True,
            },
        ]
        text = analyze_surplus(
            {
                "auction_type": "부동산임의경매",
                "claim_amount": "301,000,000원",
                "min_price": "700,000,000원",
            },
            rights,
            rights[0],
            [],
            rights_source_complete=True,
        )
        self.assertIn("선순위인 기재금액은 0원", text)
        self.assertNotIn("선순위인 기재금액은 담당자 확인 필요", text)

    def test_duplicate_expected_dividend_and_same_applicant_are_not_no_surplus_shortcuts(self):
        rights = [{
            "type": "근저당권",
            "date": "2024-01-10",
            "creditor": "A은행",
            "amount": 300_000_000,
            "isBaseRight": True,
        }]
        data = {
            "min_price": "700,000,000원",
            "auction_applicant_creditors": ["A은행"],
            "expected_dividend": {
                "auctionApplicantDividendFound": True,
                "auctionApplicantDividendAmount": 100_000_000,
            },
            "execution_cost_source_complete": True,
            "myauction_eviction_costs": {"grandTotal": 50_000_000},
        }
        text = analyze_surplus(
            data,
            rights,
            rights[0],
            [{"type": "중복경매", "caseNumber": "2024타경2"}],
            rights_source_complete=True,
        )
        self.assertIn("예상배당표상 말소기준권리 또는 경매신청채권자에게 채권배당금 100,000,000원", text)
        self.assertIn("현재 예상배당표 기준 무잉여 가능성은 없습니다", text)
        self.assertNotIn("중복경매 신청 사건이 확인되므로", text)
        self.assertNotIn("경매신청채권자는 배당을 받을 수 있으므로", text)
        self.assertNotIn("우선 배당 가능성이 높으므로", text)

        briefing_text = build_briefing_surplus_description(
            data,
            rights,
            [{"type": "중복경매", "caseNumber": "2024타경2"}],
            rights[0],
        )
        self.assertIn("현재 예상배당표 기준 무잉여 가능성은 없습니다", briefing_text)
        self.assertNotIn("70%", briefing_text)
        self.assertNotIn("중복경매 신청 사건이 확인되므로", briefing_text)

    def test_full_sale_spec_pdf_text_marks_only_that_source_complete(self):
        full_text = (
            "매각물건명세서\n최선순위 설정일자 2024.01.10\n"
            "배당요구종기 2024.03.01\n임차인 없음\n비고 해당사항 없음"
        )
        with (
            patch(
                "app.services.rights_certificate.collect_sale_spec_text_and_images",
                return_value=(full_text, ["page1.png"]),
            ),
            patch("app.services.rights_certificate.ocr_image_to_text", return_value=full_text),
        ):
            context = extract_sale_spec_tenant_context_by_ocr(object(), task_id="full-pdf")

        self.assertTrue(context["_sale_spec_complete"])
        self.assertFalse(context["_incomplete"])
        self.assertIn("매각물건명세서", context["tenant_ocr_text"])

    def test_sale_spec_header_or_empty_page_does_not_prove_complete_source(self):
        with (
            patch(
                "app.services.rights_certificate.collect_sale_spec_text_and_images",
                return_value=("매각물건명세서", ["page2.png"]),
            ),
            patch("app.services.rights_certificate.ocr_image_to_text", return_value=""),
        ):
            context = extract_sale_spec_tenant_context_by_ocr(object(), task_id="partial-pdf")

        self.assertFalse(context["_sale_spec_complete"])
        self.assertTrue(context["_incomplete"])
        self.assertIn("매각물건명세서", context["tenant_ocr_text"])

        full_text = "점유관계 임차인 없음\n비고 해당사항 없음\n최선순위 설정 2024.01.10"
        with (
            patch(
                "app.services.rights_certificate.collect_sale_spec_text_and_images",
                return_value=(full_text, ["page1.png", "page2.png"]),
            ),
            patch(
                "app.services.rights_certificate.ocr_image_to_text",
                side_effect=[full_text, "매각물건명세서"],
            ),
        ):
            mixed = extract_sale_spec_tenant_context_by_ocr(object(), task_id="mixed-pdf")

        self.assertFalse(mixed["_sale_spec_complete"])
        self.assertTrue(mixed["_incomplete"])

    def test_registry_capture_failure_is_propagated_as_incomplete(self):
        with (
            patch("app.services.rights_certificate.pytesseract", object()),
            patch(
                "app.services.rights_certificate.capturer.capture_table_split_by_rows",
                side_effect=RuntimeError("capture failed"),
            ),
        ):
            context = extract_rights_context_by_ocr(object(), task_id="capture-failure")

        self.assertTrue(context["_incomplete"])
        self.assertFalse(context["_source_complete"])
        self.assertEqual(context["rights_ocr_images"], [])

    def test_senior_lease_dividend_request_uses_three_states(self):
        right = {
            "type": "전세권",
            "date": "2023-01-01",
            "creditor": "홍길동",
            "amount": 100_000_000,
        }
        base = {"type": "근저당권", "date": "2024-01-10", "creditor": "은행"}
        self.assertIsNone(_right_dividend_request_status(right, []))
        self.assertIsNone(_right_dividend_request_status({**right, "note": "배당요구 여부 확인 필요"}, []))
        self.assertFalse(_right_dividend_request_status({**right, "note": "배당요구 없음"}, []))
        self.assertFalse(_right_dividend_request_status({**right, "note": "배당요구 미제출"}, []))
        self.assertIsNone(_right_dividend_request_status({**right, "note": "배당요구종기 2024-03-01"}, []))
        self.assertIsNone(_right_dividend_request_status({**right, "note": "배당요구 신청 예정"}, []))
        self.assertTrue(_right_dividend_request_status({
            **right,
            "note": "배당요구 접수 완료",
            "dividendRequestDate": "2024-02-01",
            "dividendDeadline": "2024-03-01",
        }, []))
        self.assertIsNone(_right_dividend_request_status(
            right,
            [{"creditor": "홍길동", "requestDate": "2024-02-01"}],
        ))
        self.assertTrue(_right_dividend_request_status(
            right,
            [{"creditor": "홍길동", "requestDate": "2024-02-01", "deadline": "2024-03-01"}],
        ))
        self.assertFalse(_right_dividend_request_status(
            right,
            [{"creditor": "홍길동", "requestDate": "2024-04-01", "deadline": "2024-03-01"}],
        ))

        unknown_text = " ".join(analyze_registered_takeover_rights([right], base, []))
        self.assertIn("배당요구 자료가 충분히 확보되지 않아", unknown_text)
        self.assertIn("인수 가능성을 배제할 수 없습니다", unknown_text)
        self.assertNotIn("낙찰자에게 인수됩니다", unknown_text)

        no_text = " ".join(analyze_registered_takeover_rights(
            [{**right, "note": "배당요구 없음"}], base, []
        ))
        self.assertIn("배당요구가 없는 것으로 확인", no_text)

        yes_text = " ".join(analyze_registered_takeover_rights(
            [right], base, [{"creditor": "홍길동", "requestDate": "2024-02-01", "deadline": "2024-03-01"}]
        ))
        self.assertIn("배당요구가 확인", yes_text)

        late_text = " ".join(analyze_registered_takeover_rights(
            [right], base, [{"creditor": "홍길동", "requestDate": "2024-04-01", "deadline": "2024-03-01"}]
        ))
        self.assertIn("배당요구 접수는 확인되지만", late_text)
        self.assertIn("접수일이 배당요구종기보다 늦어", late_text)
        self.assertNotIn("배당요구가 없는 것으로 확인", late_text)

    def test_dividend_request_parser_rejects_deadline_and_planned_rows(self):
        soup = BeautifulSoup(
            """
            <table>
              <tr><th>배당요구종기</th><td>2024-03-01</td><td>홍길동</td></tr>
              <tr><td>배당요구 신청 예정</td><td>2024-02-01</td><td>김철수</td></tr>
              <tr><td>배당요구서 접수 완료</td><td>2024-02-02</td><td>이영희</td><td>2024-03-01</td></tr>
            </table>
            """,
            "html.parser",
        )
        requests = _extract_dividend_requests(soup)
        self.assertEqual(len(requests), 1)
        self.assertEqual(requests[0]["requestDate"], "2024.02.02")
        self.assertEqual(requests[0]["deadline"], "2024.03.01")

    def test_zero_management_fee_without_explicit_clear_remains_unknown(self):
        unknown = build_checklist_from_pipeline(
            data={}, rights=[], base_right=None, valid_tenants=[],
            management_fee={"unpaidAmount": 0}, texts=[],
        )
        clear = build_checklist_from_pipeline(
            data={}, rights=[], base_right=None, valid_tenants=[],
            management_fee={"unpaidAmount": 0, "note": "미납 관리비 없음"}, texts=[],
        )
        unknown_fee = next(item for item in unknown if item.name == "체납관리비")
        clear_fee = next(item for item in clear if item.name == "체납관리비")
        self.assertEqual(unknown_fee.state, State.UNKNOWN)
        self.assertEqual(clear_fee.state, State.SAFE)

    def test_html_fallback_contains_narrative_without_table_or_internal_sources(self):
        template_path = BACKEND_PATH / "templates" / "rights_certificate" / "certificate.html"
        data = build_template_data(
            {
                "case_number": "2026타경5",
                "item_type": "아파트",
                "rights": [{"type": "근저당권", "date": "2024-01-10", "isBaseRight": True}],
                "tenants": [{"name": "임차인 없음"}],
                "management_fee": {"unpaidAmount": 0, "note": "미납 관리비 없음"},
            }
        )
        rendered = render_certificate_template(str(template_path), data)
        self.assertIn("4. 특이사항", rendered)
        self.assertIn("<ol class=\"special-list\">", rendered)
        self.assertNotIn("특이사항 종합 검토의견", rendered)
        self.assertNotIn("특이사항 종합 체크표", rendered)
        self.assertNotIn("4. 기타사항", rendered)
        self.assertNotIn("가격 분석", rendered)
        self.assertNotIn("담당자 검토", rendered)
        self.assertNotIn("관리비 채권의 소멸시효는 3년", rendered)
        self.assertNotIn("checklistRows", rendered)
        self.assertNotIn("YAML", rendered)
        self.assertNotIn("{{", rendered)

    def test_narrative_pagination_never_splits_a_block(self):
        marker = "문단중간분할금지-" + ("가" * 250)
        sections = [{
            "title": "동적 상세",
            "blocks": [
                {"heading": f"항목 {index}", "body": marker + str(index), "kind": "check", "label": "추가 확인"}
                for index in range(10)
            ],
        }]
        pages = _paginate_narrative_sections(sections)
        rendered_bodies = [block["body"] for page in pages for block in page["blocks"]]
        self.assertGreater(len(pages), 1)
        self.assertEqual(len(rendered_bodies), 10)
        for index in range(10):
            self.assertEqual(sum(body.endswith(str(index)) for body in rendered_bodies), 1)

    def test_short_narrative_blocks_fill_four_then_three_panels(self):
        sections = [{
            "title": "동적 박스",
            "blocks": [
                {"heading": f"항목 {index}", "body": "확인된 사실과 결론입니다.", "kind": "neutral", "label": "검토"}
                for index in range(7)
            ],
        }]
        pages = _paginate_narrative_sections(sections)
        self.assertEqual([len(page["blocks"]) for page in pages], [4, 3])

    def test_short_blocks_fill_pages_across_section_boundaries(self):
        sections = [
            {
                "title": "특이사항",
                "blocks": [
                    {"heading": f"요약 {index}", "body": "확인된 결론입니다.", "kind": "neutral", "label": "검토"}
                    for index in range(3)
                ],
            },
            {
                "title": "법무법인 명승 권리분석",
                "blocks": [
                    {"heading": f"상세 {index}", "body": "확인된 사실입니다.", "kind": "check", "label": "전문 검토"}
                    for index in range(4)
                ],
            },
        ]

        pages = _paginate_narrative_sections(sections)

        self.assertEqual([len(page["blocks"]) for page in pages], [4, 3])
        self.assertEqual(pages[0]["title"], "특이사항")
        self.assertEqual(
            [block["heading"] for page in pages for block in page["blocks"]],
            ["요약 0", "요약 1", "요약 2", "상세 0", "상세 1", "상세 2", "상세 3"],
        )

    def test_narrative_pagination_rebalances_a_single_panel_tail(self):
        sections = [{
            "title": "사건별 검토",
            "blocks": [
                {
                    "heading": f"항목 {index}",
                    "body": "확인된 사실과 결론을 사건자료에 따라 설명합니다. " * 6,
                    "kind": "neutral",
                    "label": "검토",
                }
                for index in range(4)
            ],
        }]

        pages = _paginate_narrative_sections(sections)

        self.assertEqual([len(page["blocks"]) for page in pages], [2, 2])

    def test_appraisal_value_is_not_used_as_unpaid_management_fee(self):
        text = "감정가 121,000,000원 미납 관리비는 관리사무소 확인 필요"

        self.assertEqual(_extract_management_fee_amount(text), 0)
        extracted = _extract_management_fee(BeautifulSoup(f"<div>{text}</div>", "html.parser"))
        self.assertEqual(extracted["unpaidAmount"], 0)

    def test_explicit_unpaid_management_fee_is_extracted(self):
        examples = (
            "미납 관리비 합계: 1,234,000원",
            "미납 관리비는 약 1,234,000원",
            "관리비 체납액 1,234,000원",
            "미납관리비 3개월분 1,500,000원",
            "미납 관리비 약 350만원",
        )

        for text in examples:
            with self.subTest(text=text):
                expected = 1_500_000 if "3개월분" in text else 3_500_000 if "350만원" in text else 1_234_000
                self.assertEqual(_extract_management_fee_amount(text), expected)
                extracted = _extract_management_fee(BeautifulSoup(f"<div>{text}</div>", "html.parser"))
                self.assertEqual(extracted["unpaidAmount"], expected)

    def test_unrelated_price_after_management_fee_keyword_is_not_used(self):
        text = "미납 관리비는 관리사무소 확인 필요. 감정가 121,000,000원"

        self.assertEqual(_extract_management_fee_amount(text), 0)


class NoTenantPlaceholderRegressionTest(unittest.TestCase):
    """다가구 명세서에서 이름을 못 읽은 '미확인' 행이 가짜 임차인으로 나열되던 회귀 방지."""

    @staticmethod
    def _rec(name="", occ="", deposit=0, rent=0, move_in="", fixed="", claim=""):
        return {
            "name": name, "occupancyType": occ, "type": occ,
            "deposit": deposit, "rent": rent,
            "moveInDate": move_in, "fixedDate": fixed, "depositClaimDate": claim,
        }

    def test_placeholder_only_records_are_treated_as_no_tenant(self):
        # 보증금·차임·일자가 전혀 없는 '미확인'/공란 레코드는 임차인 아님
        self.assertTrue(_is_no_tenant_record(self._rec(name="미확인 점유자", occ="미확인")))
        self.assertTrue(_is_no_tenant_record(self._rec(name="", occ="임차인")))

    def test_unknown_name_tenant_with_real_data_is_preserved(self):
        # 이름만 미확인이고 보증금/전입일이 있으면 실제 임차인으로 보존
        self.assertFalse(
            _is_no_tenant_record(
                self._rec(name="미확인 점유자", occ="임차인", deposit=50000000, move_in="2023-01-01")
            )
        )

    def test_no_tenant_multiunit_collapses_to_single_no_tenant_sentence(self):
        tenants = [
            self._rec(name="미확인 점유자", occ="미확인"),
            self._rec(name="미확인 점유자", occ="미확인"),
            self._rec(name="", occ="임차인"),
        ]
        text = build_tenant_analysis_text(tenants, [], "", tenant_source_complete=True)
        self.assertEqual(text, NO_TENANTS_TEXT)
        self.assertNotIn("미확인", text)

    def test_real_unknown_name_tenant_still_listed(self):
        tenants = [self._rec(name="미확인 점유자", occ="임차인", deposit=50000000, move_in="2023-01-01")]
        text = build_tenant_analysis_text(tenants, ["인수여부 확인이 필요합니다."], "", tenant_source_complete=True)
        self.assertNotEqual(text, NO_TENANTS_TEXT)
        self.assertIn("50,000,000", text)

    def test_opinion_and_certificate_consumer_renders_clean_no_tenant(self):
        # 종합의견 (2) 권리분석 → 임차권리 및 보증서 소비 경로: 무임차면 '미확인' 없이 단일 문장
        text = briefing_opinion._tenant_text_from_template_data(
            {"noTenants": True, "tenantAnalyses": [], "tenantAnalysisText": NO_TENANTS_TEXT}
        )
        self.assertNotIn("미확인", text)
        self.assertIn("임차권리는 없습니다", text)
        self.assertEqual(len([ln for ln in text.splitlines() if ln.strip()]), 1)

    def test_opinion_consumer_lists_real_tenants(self):
        text = briefing_opinion._tenant_text_from_template_data(
            {"noTenants": False, "tenantAnalyses": [
                {"description": "점유자 성명: 홍길동 / 보증금: 50,000,000원\n인수여부: 인수됩니다."},
                {"description": "점유자 성명: 김철수 / 보증금: 30,000,000원\n인수여부: 소멸됩니다."},
            ]}
        )
        self.assertIn("홍길동", text)
        self.assertIn("김철수", text)


class TenantProseFragmentRegressionTest(unittest.TestCase):
    """비고란·현황조사서 서술문에서 '대항요건을', '있고' 같은 문장 조각을 임차인명으로
    추출하고, 그 가짜 임차인이 종합의견(2) 권리분석을 2장으로 밀어내던 회귀 방지."""

    def test_name_token_rejects_sentence_fragments(self):
        for fragment in ("대항요건을", "있고", "갖추고", "하였음", "하였습니다", "되었습니다"):
            self.assertFalse(_is_name_like_token(fragment), fragment)

    def test_name_token_accepts_real_names(self):
        for name in ("홍길동", "김철수", "이영희", "박민"):
            self.assertTrue(_is_name_like_token(name), name)

    def test_guess_name_bails_on_prose_line(self):
        prose = "임차인은 대항요건을 갖추고 있고 배당요구종기 이내에 배당요구를 하였습니다."
        self.assertEqual(_guess_sale_spec_tenant_name(prose), "")

    def test_prose_remark_line_creates_no_tenant(self):
        # 날짜가 섞인 서술문도 가짜 임차인을 만들지 않는다.
        for prose in (
            "임차인은 대항요건을 갖추고 있고 배당요구를 하였습니다.",
            "임차인은 2023.01.01 전입신고를 마쳤고 대항요건을 갖추었습니다.",
        ):
            result = [t for t in parse_sale_spec_tenants_from_ocr(prose) if not _is_no_tenant_record(t)]
            self.assertEqual(result, [], prose)

    def test_labeled_tabular_rows_capture_real_tenants(self):
        ocr = (
            "점유자 성명: 홍길동 점유구분: 주거임차인 보증금: 50,000,000원 전입일: 2023.01.01 확정일: 2023.01.02 배당요구일: 2023.02.01\n"
            "점유자 성명: 김철수 점유구분: 주거임차인 보증금: 30,000,000원 전입일: 2022.03.01 확정일: 2022.03.02 배당요구일: 2022.04.01\n"
        )
        tenants = [t for t in parse_sale_spec_tenants_from_ocr(ocr) if not _is_no_tenant_record(t)]
        names = {t.get("name") for t in tenants}
        self.assertIn("홍길동", names)
        self.assertIn("김철수", names)
        self.assertNotIn("대항요건을", names)

    def test_no_tenant_opinion_fits_single_page(self):
        opinion = "\n".join([
            "1) 말소기준 및 등기부상 소멸사항",
            "- 2023.05.16 설정된 근저당권[○○대부]이 말소기준권리이며 이후 권리는 모두 소멸되어 인수하는 권리는 없습니다.",
            "2) 임차권리 인수사항",
            "- 조사된 임차인이 없으므로, 낙찰자에게 인수되는 임차권리는 없습니다.",
            "3) 경매취하 / 무잉여 가능성",
            "- 경매신청 채권자의 청구금액은 761,315,340원입니다.",
            "- 등기부상 채권 총액이 최저가보다 높아 취하 가능성은 낮습니다.",
            "- 신청채권자는 경매비용·당해세 다음으로 배당받을 수 있어 무잉여 가능성은 낮습니다.",
        ])
        self.assertEqual(len(ppt_builder._split_rights_analysis_opinion_pages(opinion)), 1)

    def test_many_tenants_opinion_splits_to_two_pages(self):
        lines = ["1) 말소기준 및 등기부상 소멸사항",
                 "- 2023.05.16 설정된 근저당권[○○대부]이 말소기준권리이며 이후 권리는 모두 소멸됩니다.",
                 "2) 임차권리 인수사항"]
        for i in range(6):
            lines.append(
                f"점유자 성명: 임차인{i} / 점유구분: 주거임차인 / 보증금: 50,000,000원 / 차임: 없음 / 전입일: 2023.01.0{i} / 확정일: 2023.01.0{i} / 배당요구일: 2023.02.0{i}"
            )
            lines.append("인수여부: 대항력 및 우선변제권 성립 여부는 배당요구종기일 기준으로 최종 검토가 필요합니다.")
        lines += ["3) 경매취하 / 무잉여 가능성",
                  "- 경매신청 채권자의 청구금액은 761,315,340원입니다."]
        self.assertGreaterEqual(len(ppt_builder._split_rights_analysis_opinion_pages("\n".join(lines))), 2)

    def test_customer_safe_text_strips_staff_confirm_token(self):
        cleaned = briefing_rights._customer_safe_text("보증금: 담당자 확인 필요 / 차임: 없음")
        self.assertNotIn("담당자 확인 필요", cleaned)
        self.assertIn("미확인", cleaned)

    def test_linter_flags_staff_confirm_token(self):
        issues = briefing_opinion.lint_customer_facing_text("보증금: 담당자 확인 필요", "rights")
        self.assertTrue(any("staff-confirm-token" in issue for issue in issues))

    # --- 매각물건명세서 PDF텍스트 경로 (실제 운영 1차 파서) ---
    def test_sale_spec_name_before_source_rejects_fragments(self):
        self.assertEqual(_sale_spec_name_before_source("임차인은 대항요건을 갖추고 있고"), "")
        self.assertEqual(_sale_spec_name_before_source("홍길동 전부"), "홍길동")

    def test_pdf_text_prose_rows_create_no_fake_tenants(self):
        # 비고/각주 서술문이 점유자 '표의 행'으로 오인되어 조각이 임차인이 되지 않는다.
        text = "\n".join([
            "점유자 성명 점유부분 정보출처 점유의 권원 전입신고일자 확정일자 (배당요구일자)",
            "홍길동 전부 현황조사 주거임차인 50,000,000 2023.01.01 2023.01.02 2023.02.01",
            "김철수 전부 권리신고 주거임차인 30,000,000 2022.03.01 2022.03.02 2022.04.01",
            "현황조사서상 임차인은 대항요건을 갖추고 있고 배당요구종기 이내에 배당요구를 하였습니다",
        ])
        tenants = [t for t in parse_sale_spec_tenants_from_pdf_text(text) if not _is_no_tenant_record(t)]
        names = {t.get("name") for t in tenants}
        self.assertIn("홍길동", names)
        self.assertIn("김철수", names)
        self.assertNotIn("대항요건을", names)
        self.assertNotIn("있고", names)

    def test_sale_spec_metadata_lines_do_not_create_fake_tenant(self):
        # '조사된 임차내역없음'인데 상단 메타데이터('최선순위 설정 … 근저당권 배당요구종기 …')에서
        # '설정'·'종기'를 임차인명으로, 설정일자를 전입일로 뽑던 실제 사건(서울서부 2024타경52693) 회귀.
        text = "\n".join([
            "최저매각가격의 표시 별지 기재와 같음 최선순위",
            "설정 2018.6.11 근저당권 배당요구종기 2024. 5. 27.",
            "점유자 성명 점유부분 정보출처 ... 확정일자 배당요구여부(배당요구일자)",
            "조사된 임차내역없음",
        ])
        for parser in (parse_sale_spec_tenants_from_ocr, parse_sale_spec_tenants_from_pdf_text):
            valid = [t for t in parser(text) if not _is_no_tenant_record(t)]
            self.assertEqual(valid, [], f"{parser.__name__} produced {[t.get('name') for t in valid]}")

    def test_select_best_prefers_occupancy_region_over_ocr_fakes(self):
        # 실제 사건(서울서부 2024타경, 권미선 1명)에서 명세서 셀이 한 줄씩 쪼개져 pdf_text 파서가
        # summary_fallback 으로 권미선만 읽었는데, 전문서 OCR이 '있는'·'변제' 가짜를 더해 개수로
        # 이기던 회귀 방지 — 점유자 영역에서 읽은 결과를 전문서 OCR보다 신뢰해야 한다.
        from app.services.rights_certificate import _select_best_sale_spec_tenants
        occupancy = [{
            "name": "권미선", "occupancyType": "주거 임차인", "type": "주거 임차인",
            "deposit": 190000000, "rent": 0,
            "moveInDate": "2022.08.22", "fixedDate": "", "depositClaimDate": "",
            "_parse_method": "summary_fallback",
        }]
        ocr_with_fakes = [
            {"name": "권미선", "deposit": 0, "moveInDate": "2024.08.20"},
            {"name": "있는", "deposit": 190000000, "moveInDate": "2022.08.22"},
            {"name": "변제", "deposit": 0, "moveInDate": "2022.07.25"},
        ]
        best = _select_best_sale_spec_tenants(occupancy, ocr_with_fakes)
        names = {t.get("name") for t in best if not _is_no_tenant_record(t)}
        self.assertEqual(names, {"권미선"})


class MultiTenantBlockParseTest(unittest.TestCase):
    """세로로 쪼개진 점유자 표(다가구·상가 다수임차인)를 임차인별로 재구성하는지 검증.
    실측: 서울남부 2026타경114(상가 4명), 서울남부 2024타경105895(다가구 7명)."""

    FIX_114 = "\n".join([
        "김호은 현황조사 - 임차인 2026.02.20.",
        "장만희",
        "1층 현황조사 점포",
        "임차인 2022.07.11.",
        "1층 권리신고 점포",
        "임차인",
        "2022.08.10.", "부터", "2028.08.10.", "까지",
        "100,000,000 3,500,000 2022.07.11 2026.04.09 2026.4.17.",
        "주식회", "사", "컴타운",
        "3층 현황조사 점포",
        "임차인 2021.06.09.",
        "3층", "전부 권리신고 점포", "임차인",
        "2022.10.17.", "부터", "2026.10.16.", "까지",
        "10,000,000 1,000,000 2022.10.26 2026.4.3.",
        "하나협", "동조합 2층 현황조사 점포",
        "임차인 미상",
    ]).splitlines()

    def test_114_four_tenants_split_with_amounts(self):
        tenants = _parse_sale_spec_occupancy_blocks(self.FIX_114)
        names = [t["name"] for t in tenants]
        self.assertEqual(names, ["김호은", "장만희", "주식회사 컴타운", "하나협동조합"])
        by = {t["name"]: t for t in tenants}
        self.assertEqual(by["장만희"]["deposit"], 100000000)
        self.assertEqual(by["장만희"]["rent"], 3500000)
        self.assertEqual(by["장만희"]["moveInDate"], "2022.07.11")
        self.assertEqual(by["주식회사 컴타운"]["deposit"], 10000000)
        self.assertEqual(by["주식회사 컴타운"]["moveInDate"], "2022.10.26")

    def test_block_parser_rejects_names_ending_in_ho(self):
        # '안영호'처럼 '호'로 끝나는 이름을 점유부분(301호)으로 오판하지 않는다.
        lines = ["안영호", "2층", "301호", "현황조사", "주거", "임차인", "2024.01.03"]
        tenants = _parse_sale_spec_occupancy_blocks(lines)
        self.assertEqual([t["name"] for t in tenants], ["안영호"])

    def test_daehangnyeok_marked_per_tenant(self):
        base = {"date": "2018.06.20", "type": "근저당권"}
        # 전입이 말소기준보다 늦음 → 대항력 X
        late = [{"name": "장만희", "moveInDate": "2022.07.11", "fixedDate": "2026.04.09", "depositClaimDate": "2026.04.17", "deposit": 100000000}]
        texts = analyze_tenants(late, base, [], "2026.06.12", "서울")
        self.assertTrue(any("대항력 X" in t and "인수되는 임차권리는 없습니다" in t for t in texts))
        # 전입이 말소기준보다 빠름 → 대항력 O
        early = [{"name": "김선순", "moveInDate": "2017.01.01", "fixedDate": "2017.01.02", "depositClaimDate": "2018.01.01", "deposit": 100000000}]
        texts2 = analyze_tenants(early, base, [], "2018.06.12", "서울")
        self.assertTrue(any("대항력 O" in t for t in texts2))

    def test_analyze_tenants_not_deduped_for_multi(self):
        base = {"date": "2018.06.20"}
        tenants = [
            {"name": "A", "moveInDate": "2022.01.01", "deposit": 0},
            {"name": "B", "moveInDate": "2023.01.01", "deposit": 0},
            {"name": "C", "moveInDate": "2024.01.01", "deposit": 0},
        ]
        texts = analyze_tenants(tenants, base, [], "", "")
        self.assertEqual(len(texts), 3)


class RightsOpinionPaginationTest(unittest.TestCase):
    """담당자 종합의견 (2) 권리분석이 길어 추가 페이지가 생길 때 본문 텍스트 박스가
    중복 생성(겹침)되지 않는지 검증."""

    def _body_boxes(self, slide):
        boxes = []
        for sh in slide.shapes:
            if not getattr(sh, "has_text_frame", False):
                continue
            if (sh.text or "").strip().isdigit():
                continue
            if int(sh.width) * int(sh.height) > Inches(2) * Inches(2):
                boxes.append(sh)
        return boxes

    def test_duplicate_slide_does_not_double_placeholders(self):
        from app.core.config import settings
        prs = Presentation(settings.pptm_template)
        slide = ppt_builder.find_slide_by_note_key(prs, "SLIDE_KEY=OPINION_RIGHTS_ANALYSIS")
        if slide is None:
            self.skipTest("권리분석 템플릿 슬라이드 없음")
        before = sum(1 for sh in slide.shapes if getattr(sh, "has_text_frame", False))
        dup = ppt_builder.duplicate_slide(prs, slide)
        after = sum(1 for sh in dup.shapes if getattr(sh, "has_text_frame", False))
        self.assertEqual(after, before)

    def test_extra_page_has_single_body_box(self):
        from app.core.config import settings
        prs = Presentation(settings.pptm_template)
        slide = ppt_builder.find_slide_by_note_key(prs, "SLIDE_KEY=OPINION_RIGHTS_ANALYSIS")
        if slide is None:
            self.skipTest("권리분석 템플릿 슬라이드 없음")
        base = prs.slides.index(slide)
        long_text = "\n".join(
            [f"{n}) 섹션" for n in (1, 2, 3)] + ["- " + ("가" * 60) for _ in range(20)]
        )
        pages = ppt_builder._split_rights_analysis_opinion_pages(long_text)
        self.assertGreaterEqual(len(pages), 2)
        ppt_builder.apply_rights_analysis_opinion(prs, long_text)
        for i in range(base, base + len(pages)):
            self.assertLessEqual(len(self._body_boxes(prs.slides[i])), 1, f"slide {i} 본문박스 중복")


if __name__ == "__main__":
    unittest.main()
